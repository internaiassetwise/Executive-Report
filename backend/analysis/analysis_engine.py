"""Domain-agnostic deterministic analysis. Runs unchanged in CPython and Pyodide.
No uploaded content is executed. Original cells remain in the worker-owned store.
"""
import csv, io, json, math, re, statistics, zipfile
from datetime import date, datetime
from collections import Counter, defaultdict

TABLES = {}
WORKBOOK = {}
MAX_CELLS = 400_000
MAX_ROWS = 50_000

def present(v):
    return v is not None and (not isinstance(v, str) or bool(v.strip()))

def number(v):
    if isinstance(v, bool) or isinstance(v, (date, datetime)):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if not isinstance(v, str):
        return None
    s = v.strip()
    if re.match(r'^0\d+', s):
        return None
    if re.fullmatch(r'[+-]?\d+(\.\d+)?([eE][+-]?\d+)?', s):
        n = float(s)
        return n if math.isfinite(n) else None
    return None

def iso(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()[:10]
    if isinstance(v, str) and re.match(r'^\d{4}-\d{2}-\d{2}($|T)', v):
        try:
            return date.fromisoformat(v[:10]).isoformat()
        except ValueError:
            pass
    return None

def clean(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v.strip() if isinstance(v, str) else v

def fmt(n):
    if n is None:
        return '—'
    return f'{n:,.2f}'.rstrip('0').rstrip('.') if isinstance(n, (int, float)) else str(n)

def col_letter(i):
    out = ''
    while i:
        i, r = divmod(i - 1, 26)
        out = chr(65 + r) + out
    return out

def quantile(values, p):
    vals = sorted(values)
    pos = (len(vals) - 1) * p
    a = int(pos)
    b = min(a + 1, len(vals) - 1)
    return vals[a] + (vals[b] - vals[a]) * (pos - a)

def blocks(indices, gap=1):
    result = []
    for x in sorted(indices):
        if not result or x - result[-1][-1] > gap:
            result.append([x])
        else:
            result[-1].append(x)
    return result

def load_sheets(raw, filename):
    if filename.lower().endswith('.csv'):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp874')
        try:
            dialect = csv.Sniffer().sniff(text[:16000], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        grid = list(csv.reader(io.StringIO(text), dialect))
        return [{'name': 'CSV', 'grid': grid, 'merges': [], 'hidden_rows': [], 'hidden_columns': [], 'formulas': [], 'errors': [], 'state': 'visible', 'formats': {}}]
    if filename.lower().endswith('.xls'):
        import xlrd
        book = xlrd.open_workbook(file_contents=raw, formatting_info=True)
        sheets = []
        for sh in book.sheets():
            if sh.nrows * sh.ncols > MAX_CELLS:
                raise ValueError('ตารางใหญ่เกินขีดจำกัด 400,000 เซลล์ กรุณาแบ่งไฟล์')
            grid = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    row.append(value)
                grid.append(row)
            sheets.append({'name': sh.name, 'grid': grid, 'merges': [(a+1,c+1,b,d) for a,b,c,d in sh.merged_cells], 'hidden_rows': [r+1 for r,v in sh.rowinfo_map.items() if v.hidden], 'hidden_columns': [col_letter(c+1) for c,v in sh.colinfo_map.items() if v.hidden], 'formulas': [], 'errors': [], 'state': {0:'visible',1:'hidden',2:'veryHidden'}.get(getattr(sh,'visibility',0),'hidden'), 'formats': {}, 'legacy': True})
        return sheets
    import openpyxl
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(x.file_size for x in z.infolist()) > 60_000_000:
            raise ValueError('ข้อมูลหลังคลายไฟล์มีขนาดเกิน 60 MB กรุณาแบ่งไฟล์')
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
    sheets = []
    total_cells = 0
    for sh in wb:
        # Read only material cells; formatted empty tails are ignored.
        material = [c for c in sh._cells.values() if present(c.value)]
        if not material:
            sheets.append({'name': sh.title, 'grid': [], 'merges': [], 'hidden_rows': [], 'hidden_columns': [], 'formulas': [], 'errors': [], 'state': sh.sheet_state, 'formats': {}})
            continue
        mr, mc = max(c.row for c in material), max(c.column for c in material)
        total_cells += mr * mc
        if total_cells > MAX_CELLS or mr > MAX_ROWS:
            raise ValueError('ไฟล์เกินขีดจำกัด 400,000 เซลล์ หรือ 50,000 แถว กรุณาแบ่งไฟล์')
        grid = [[None] * mc for _ in range(mr)]
        formulas, errors, formats = [], [], {}
        for cell in material:
            v = cell.value
            formats[f'{cell.row},{cell.column}'] = cell.number_format
            if cell.data_type == 'f':
                v = cached[sh.title][cell.coordinate].value
                formulas.append({'cell': cell.coordinate, 'formula': cell.value, 'cached': v is not None})
            elif cell.data_type == 'e':
                errors.append({'cell': cell.coordinate, 'value': v})
                v = None
            grid[cell.row-1][cell.column-1] = v
        sheets.append({'name':sh.title,'grid':grid,'merges':[(m.min_row,m.min_col,m.max_row,m.max_col) for m in sh.merged_cells.ranges], 'hidden_rows':[r for r,v in sh.row_dimensions.items() if v.hidden], 'hidden_columns':[c for c,v in sh.column_dimensions.items() if v.hidden], 'formulas':formulas,'errors':errors,'state':sh.sheet_state,'formats':formats})
    return sheets

def profile(name, vals, formats):
    nonempty = [v for v in vals if present(v)]
    ns = [number(v) for v in nonempty]
    nums = [n for n in ns if n is not None]
    dates = [iso(v) for v in nonempty]
    unique = len(set(str(clean(v)) for v in nonempty))
    n = len(nonempty)
    numeric = bool(n) and len(nums) / n >= .85
    is_date = bool(n) and sum(v is not None for v in dates) / n >= .9
    id_name = bool(re.search(r'(^id$|\bid\b|รหัส|code|sku|เลขที่|identifier)',name,re.I))
    codes = bool(n) and all(isinstance(v,str) and bool(re.fullmatch(r'(?:0\d+|[A-Za-z]+[-_]\d+)',v.strip())) for v in nonempty)
    is_id = id_name or codes
    dtype = 'datetime' if is_date else 'numeric' if numeric else 'boolean' if n and all(isinstance(v,bool) for v in nonempty) else 'text'
    role = 'identifier' if is_id else 'time_dimension' if is_date else 'measure' if numeric else 'dimension' if unique <= 30 else 'label'
    semantic = 'identifier' if is_id else 'datetime' if is_date else 'percentage' if any('%' in f for f in formats) else 'numeric_measure' if numeric else 'category' if unique <= 30 else 'free_text'
    result = {'name':name,'normalized_name':re.sub(r'\s+','_',name.strip().lower()),'type':dtype,'semantic_type':semantic,'role':role,'unique':unique,'missing':len(vals)-n,'null_pct':round((len(vals)-n)/len(vals)*100,2) if vals else 0,'valid_count':len(nums) if numeric else n,'invalid_count':n-len(nums) if numeric else 0,'samples':[str(clean(v))[:100] for v in nonempty[:3]],'confidence':.9 if n and (numeric or is_date or is_id) else .7}
    result['numeric_unique'] = len(set(nums))
    if nums and numeric:
        result['stats']={'count':len(nums),'mean':statistics.mean(nums),'median':statistics.median(nums),'min':min(nums),'max':max(nums),'std':statistics.stdev(nums) if len(nums)>1 else 0}
    return result

def detect(sheet):
    grid = sheet['grid']
    if not grid:
        return []
    if len(grid)>MAX_ROWS or sum(len(r) for r in grid)>MAX_CELLS:
        raise ValueError('ไฟล์เกินขีดจำกัด 400,000 เซลล์ หรือ 50,000 แถว')
    width = max(len(row) for row in grid)
    grid = [row + [None]*(width-len(row)) for row in grid]
    # Ignore isolated metadata for segmentation. Keep it in workbook profile.
    active = [r for r,row in enumerate(grid) if any(present(v) for v in row)]
    found=[]
    bands=[]
    for group in blocks(active, gap=2):
        current=[]
        for r in group:
            filled=[v for v in grid[r] if present(v)]
            new_header=len(filled)>=2 and all(isinstance(v,str) and number(v) is None and iso(v) is None for v in filled)
            if current and r-current[-1]>1 and new_header:
                bands.append(current);current=[]
            current.append(r)
        if current:bands.append(current)
    for band in bands:
        if len(band)<2:
            continue
        populated_cols=[c for c in range(width) if any(present(grid[r][c]) for r in band)]
        for cb in blocks(populated_cols):
            lo,hi=cb[0],cb[-1]+1
            if len(cb)<1:
                continue
            candidates=[]
            for r in band[:15]:
                cells=grid[r][lo:hi]
                filled=[v for v in cells if present(v)]
                text_count=sum(isinstance(v,str) and number(v) is None and iso(v) is None for v in filled)
                unique=len(set(str(v) for v in filled))
                contrast=0
                for c,v in enumerate(cells):
                    later=[grid[k][lo+c] for k in band if r<k<=r+6 and present(grid[k][lo+c])]
                    if isinstance(v,str) and number(v) is None and iso(v) is None and later and sum(number(x) is not None or iso(x) is not None for x in later)/len(later)>=.75:contrast+=1
                score=(text_count/max(1,len(cells)))*.55+(len(filled)/len(cells))*.25+(unique/max(1,len(cells)))*.2+contrast/max(1,len(cells))*.5
                if text_count/max(1,len(filled))<.8 and not contrast:score=0
                if all(number(v) is not None for v in filled): score=0
                candidates.append((score,r))
            best=max(candidates,key=lambda x:(x[0],-x[1]))
            header=best[1] if best[0]>=.65 else None
            start=(header+1) if header is not None else band[0]
            names=[]
            used=Counter()
            for c in range(lo,hi):
                base=str(grid[header][c]).strip() if header is not None and present(grid[header][c]) else f'Column {c+1}'
                for r1,c1,r2,c2 in sheet['merges']:
                    if header is not None and r2 == header and c1<=c+1<=c2 and present(grid[r1-1][c1-1]):
                        base=f'{grid[r1-1][c1-1]} / {base}'
                used[base]+=1
                names.append(base if used[base]==1 else f'{base} ({used[base]})')
            rows=[]; original=[]; source_rows=[]; excluded=[]; segment_start=0; unresolved_summaries=[]
            for r in range(start,band[-1]+1):
                vals=grid[r][lo:hi]
                if not any(present(v) for v in vals):
                    excluded.append({'row':r+1,'reason':'แถวว่าง'});continue
                # Only exclude a labeled summary if a numeric cell reconciles to preceding rows.
                label=next((str(v).strip().lower() for v in vals if present(v)), '')
                is_total=bool(re.fullmatch(r'(grand total|total|subtotal|รวม|รวมทั้งหมด|ยอดรวม|รวมย่อย)',label))
                reconciles=False
                if is_total and rows:
                    for c,v in enumerate(vals):
                        nv=number(v)
                        for preceding_rows in (rows[segment_start:],rows):
                            nums=[number(row[c]) for row in preceding_rows if number(row[c]) is not None]
                            if nv is not None and len(nums)>=2 and math.isclose(math.fsum(nums),nv,rel_tol=1e-7,abs_tol=.005):reconciles=True
                if reconciles:
                    excluded.append({'row':r+1,'reason':'แถวสรุปยอดที่ตรวจสอบกับผลรวมแถวก่อนหน้าแล้ว'});segment_start=len(rows);continue
                if is_total:unresolved_summaries.append(r+1)
                original.append(vals);rows.append([clean(v) for v in vals]);source_rows.append(r+1)
            if not rows:continue
            columns=[profile(name,[r[i] for r in original],[sheet['formats'].get(f'{r},{lo+i+1}','') for r in source_rows]) for i,name in enumerate(names)]
            issues=[]
            if unresolved_summaries:issues.append({'kind':'summary','severity':'warning','count':len(unresolved_summaries),'message':f'พบแถวที่อาจเป็นยอดรวมแต่ตรวจสอบไม่ได้ที่แถว {unresolved_summaries} ยังรวมอยู่ในผล กรุณาตรวจและแก้ไขไฟล์ก่อนใช้ตัวเลข'})
            missing=sum(c['missing'] for c in columns)
            duplicate=len(rows)-len(set(json.dumps(r,ensure_ascii=False,default=str) for r in rows))
            if missing:issues.append({'kind':'missing','severity':'warning','count':missing,'message':f'พบค่าว่าง {fmt(missing)} เซลล์ เก็บเป็นค่าว่างและไม่นำไปคำนวณสถิติของคอลัมน์นั้น'})
            if duplicate:issues.append({'kind':'duplicates','severity':'warning','count':duplicate,'message':f'พบแถวซ้ำ {fmt(duplicate)} แถว ยังคงรวมอยู่ในการวิเคราะห์'})
            if excluded:issues.append({'kind':'excluded','severity':'info','count':len(excluded),'message':f'ไม่นำแถวว่างหรือสรุปยอด {fmt(len(excluded))} แถวมาคำนวณ ดูรายการแถวที่แยกออก'})
            for c in columns:
                if c['invalid_count']:issues.append({'kind':'mixed','severity':'warning','count':c['invalid_count'],'message':f"{c['name']}: มีค่าที่แปลงเป็นตัวเลขไม่ได้ {c['invalid_count']} ค่า"})
                if c['unique']==1:issues.append({'kind':'constant','severity':'info','count':1,'message':f"{c['name']}: มีค่าเดียว ไม่เหมาะกับการวิเคราะห์ความสัมพันธ์"})
                if c['role']=='label':issues.append({'kind':'cardinality','severity':'info','count':c['unique'],'message':f"{c['name']}: มีค่าต่างกันจำนวนมาก ไม่สร้างกราฟแยกหมวดหมู่อัตโนมัติ"})
            ambiguities=sum(isinstance(v,str) and bool(re.fullmatch(r'\d{1,2}/\d{1,2}/\d{2,4}',v.strip())) for row in original for v in row)
            if ambiguities:issues.append({'kind':'dates','severity':'warning','count':ambiguities,'message':'พบวันที่แบบมีเครื่องหมาย / ที่ยังไม่ทราบรูปแบบ เก็บเป็นข้อความ กรุณาใช้ YYYY-MM-DD หากต้องการวิเคราะห์แนวโน้ม'})
            table_id=f'T{len(TABLES)+1:03d}'
            if header is not None and header>band[0]:issues.append({'kind':'metadata','severity':'info','count':header-band[0],'message':f'แถว {band[0]+1}–{header} ก่อนหัวตารางเก็บเป็นข้อมูลประกอบ ไม่รวมในผลคำนวณ'})
            t={'id':table_id,'name':f"{sheet['name']} · ตาราง {len(found)+1}",'sheet':sheet['name'],'range':f'{col_letter(lo+1)}{(header if header is not None else start)+1}:{col_letter(hi)}{band[-1]+1}','header_row':header+1 if header is not None else None,'confidence':round(min(1,best[0]),2) if header is not None else .45,'rows_count':len(rows),'columns_count':len(columns),'columns':columns,'preview':rows[:12],'quality':{'missing':missing,'duplicates':duplicate,'completeness':round(100-missing/(len(rows)*len(columns))*100,1),'issues':issues},'excluded_rows':excluded,'source_rows':source_rows,'column_offset':lo,'rows':rows,'original':original}
            TABLES[table_id]=t
            t['opportunities']=plan(t)
            found.append({k:v for k,v in t.items() if k not in ('rows','original','source_rows')})
    return found

def plan(t):
    cols=t['columns']; plans=[]
    measures=[i for i,c in enumerate(cols) if c['role']=='measure']
    dims=[i for i,c in enumerate(cols) if c['role']=='dimension' and 2<=c['unique']<=30]
    dates=[i for i,c in enumerate(cols) if c['role']=='time_dimension' and c['unique']>=3]
    def add(kind,title,reason,indices):
        plans.append({'id':f'{t["id"]}-{len(plans)+1:02d}','type':kind,'title':title,'reason':reason,'columns':indices,'source_columns':[cols[i]['name'] for i in indices]})
    add('quality','คุณภาพข้อมูล','ตรวจค่าว่างและแถวซ้ำ โดยเก็บข้อมูลต้นฉบับไว้',list(range(len(cols))))
    for i in measures:
        c=cols[i]; add('statistics',f'สถิติสรุป · {c["name"]}','ค่าเฉลี่ย มัธยฐาน ช่วงค่า และส่วนเบี่ยงเบนมาตรฐาน',[i])
        if c['valid_count']>=8 and c['numeric_unique']>=3:add('outliers',f'ค่าผิดปกติ · {c["name"]}','ตรวจค่าที่อยู่นอกช่วง 1.5 × IQR ไม่ถือว่าเป็นข้อผิดพลาดอัตโนมัติ',[i])
        if c['valid_count']>=10 and c['numeric_unique']>=3:add('distribution',f'การกระจาย · {c["name"]}','จัดกลุ่มความถี่ของค่าตัวเลขเพื่อดูรูปแบบการกระจาย',[i])
    for i in dims[:4]:
        add('frequency',f'สัดส่วนหมวดหมู่ · {cols[i]["name"]}','เปรียบเทียบจำนวนแถวในแต่ละหมวดหมู่',[i])
        for j in measures[:2]:add('category',f'{cols[j]["name"]} ตาม {cols[i]["name"]}','เปรียบเทียบค่าเฉลี่ยต่อหมวดหมู่ ไม่สมมติว่าตัวเลขบวกรวมกันได้',[i,j])
    for i in dates[:1]:
        for j in measures[:3]:add('trend',f'แนวโน้ม · {cols[j]["name"]}','ค่าเฉลี่ยรายเดือน ไม่เติมช่วงเวลาที่ไม่มีข้อมูล',[i,j])
    for a,i in enumerate(measures[:5]):
        for j in measures[:5][a+1:]:
            pairs=[(number(r[i]),number(r[j])) for r in t['rows'] if number(r[i]) is not None and number(r[j]) is not None]
            if len(pairs)>=10 and len(set(x for x,y in pairs))>1 and len(set(y for x,y in pairs))>1:
                add('correlation',f'{cols[i]["name"]} ↔ {cols[j]["name"]}','Pearson correlation ของแถวที่มีค่าครบ ไม่สรุปเหตุและผล',[i,j])
    return plans

def inspect(raw, filename):
    TABLES.clear()
    WORKBOOK.clear()
    if not raw:raise ValueError('ไฟล์ว่าง กรุณาเลือกไฟล์ที่มีข้อมูล')
    if len(raw)>15*1024*1024:raise ValueError('ไฟล์ใหญ่กว่า 15 MB กรุณาแบ่งไฟล์')
    sheets=load_sheets(raw,filename)
    tables=[]; notes=[]
    for sheet in sheets:
        tables.extend(detect(sheet))
        if sheet['hidden_rows'] or sheet['hidden_columns'] or sheet['state']!='visible':notes.append(f"{sheet['name']}: รวมข้อมูลจากแถว คอลัมน์ หรือชีตที่ซ่อนอยู่แล้ว")
        uncached=sum(not f['cached'] for f in sheet['formulas'])
        if uncached:notes.append(f"{sheet['name']}: สูตร {uncached} เซลล์ไม่มีค่าที่คำนวณไว้ ให้เปิดไฟล์ใน Excel แล้วบันทึกใหม่ก่อนวิเคราะห์")
        if sheet['errors']:notes.append(f"{sheet['name']}: ข้อผิดพลาดสูตร {len(sheet['errors'])} เซลล์ ถูกเก็บเป็นค่าว่าง")
        if sheet.get('legacy'):notes.append(f"{sheet['name']}: ไฟล์ XLS ใช้ค่าที่บันทึกไว้ ไม่สามารถตรวจสูตรเดิมได้ครบถ้วน")
    if not tables:raise ValueError('ไม่พบตารางที่มีอย่างน้อยสองแถว กรุณาตรวจโครงสร้างไฟล์')
    if any(t['confidence']<.7 for t in tables):notes.append('บางตารางมีความมั่นใจในการหาหัวตารางต่ำ กรุณาตรวจตัวอย่างข้อมูลก่อนวิเคราะห์')
    coverage=[]
    for sheet in sheets:
        members=[t for t in tables if t['sheet']==sheet['name']]
        coverage.append({'name':sheet['name'],'state':sheet['state'],'tables_count':len(members),'rows_count':sum(t['rows_count'] for t in members),'table_ids':[t['id'] for t in members],'status':'ready' if members else 'no_table','reason':'' if members else ('ชีตว่าง' if not sheet['grid'] else 'ไม่พบตารางที่มีอย่างน้อยสองแถว')})
    titles={'quality':'คุณภาพข้อมูลทุกตาราง','statistics':'สถิติเบื้องต้น','outliers':'ค่าผิดปกติ','distribution':'การกระจายของตัวเลข','frequency':'สัดส่วนหมวดหมู่','category':'เปรียบเทียบตามหมวดหมู่','trend':'แนวโน้มตามเวลา','correlation':'ความสัมพันธ์ระหว่างตัวเลข'}
    opportunities=[]
    for kind,title in titles.items():
        plans=[p for t in tables for p in t['opportunities'] if p['type']==kind]
        if plans:opportunities.append({'type':kind,'title':title,'reason':plans[0]['reason'],'analyses_count':len(plans),'tables_count':sum(any(p['type']==kind for p in t['opportunities']) for t in tables)})
    profile={'filename':filename,'sheets_count':len(sheets),'tables_count':len(tables),'rows_count':sum(t['rows_count'] for t in tables),'tables':tables,'sheets':coverage,'summary':summarize_tables(tables),'opportunities':opportunities,'notes':notes,'workbook_profile':[{'name':s['name'],'merged_ranges':s['merges'],'hidden_rows':s['hidden_rows'],'hidden_columns':s['hidden_columns'],'formulas':s['formulas'],'errors':s['errors']} for s in sheets],'understanding':{'dataset_summary':f'อ่านครบ {len(sheets)} ชีต · พบ {len(tables)} ตาราง พร้อมวิเคราะห์ทั้งหมดอัตโนมัติ','possible_domain':'ไม่กำหนดประเภทธุรกิจ','grain':'คำนวณแยกแต่ละตาราง แล้วรวมข้อค้นพบในรายงานเดียวพร้อมระบุชีตต้นทาง','limitations':['การแยกตารางและบทบาทคอลัมน์เป็นการอนุมาน ควรตรวจทานก่อนใช้ตัดสินใจ','ไม่คำนวณสูตร Excel ใหม่ และไม่อนุมานหน่วย เงินสกุล หรือเหตุและผล']}}
    WORKBOOK.update(profile)
    return profile

def table_source(t):
    return {'table_id':t['id'],'sheet':t['sheet'],'range':t['range'],'source_columns':[]}

def summarize_tables(tables):
    cells=sum(t['rows_count']*t['columns_count'] for t in tables)
    missing=sum(t['quality']['missing'] for t in tables)
    return {'rows_count':sum(t['rows_count'] for t in tables),'columns_count':sum(t['columns_count'] for t in tables),'measures_count':sum(c['role']=='measure' for t in tables for c in t['columns']),'cells_count':cells,'quality':{'missing':missing,'duplicates':sum(t['quality']['duplicates'] for t in tables),'completeness':round((cells-missing)/cells*100,1) if cells else 0,'issues':[{**issue,'source':table_source(t)} for t in tables for issue in t['quality']['issues']]}}

def analyze(table_id, selected, objective=''):
    if table_id not in TABLES:raise ValueError('ไม่พบชุดข้อมูล กรุณาอัปโหลดใหม่')
    t=TABLES[table_id]; cols=t['columns'];rows=t['rows']; results=[]; evidence=[]
    allowed={p['id']:p for p in t['opportunities']}
    if not selected or any(s not in allowed for s in selected):raise ValueError('กรุณาเลือกการวิเคราะห์ที่ใช้ได้อย่างน้อยหนึ่งรายการ')
    for sid in dict.fromkeys(selected):
        p=allowed[sid];kind=p['type'];idx=p['columns']; data={}; chart=None; finding=''
        if kind=='quality':
            data={k:t['quality'][k] for k in ['completeness','missing','duplicates']}
            finding=f"ข้อมูลครบถ้วน {fmt(data['completeness'])}% พบค่าว่าง {fmt(data['missing'])} เซลล์ และแถวซ้ำ {fmt(data['duplicates'])} แถว"
        elif kind=='statistics':
            data=cols[idx[0]]['stats'];finding=f"{cols[idx[0]]['name']} มีค่าเฉลี่ย {fmt(data['mean'])} มัธยฐาน {fmt(data['median'])} จากข้อมูล {fmt(data['count'])} ค่า"
        elif kind=='outliers':
            vals=[number(r[idx[0]]) for r in rows if number(r[idx[0]]) is not None];q1=quantile(vals,.25);q3=quantile(vals,.75);iqr=q3-q1
            low=q1-1.5*iqr;high=q3+1.5*iqr
            points=[{'row':t['source_rows'][i],'value':number(r[idx[0]])} for i,r in enumerate(rows) if number(r[idx[0]]) is not None and (number(r[idx[0]])<low or number(r[idx[0]])>high)] if iqr>0 else []
            data={'count':len(points),'lower':low,'upper':high,'iqr':iqr,'points':points,'sample_size':len(vals)}
            finding=f"พบ {fmt(len(points))} ค่านอกช่วง {fmt(low)} ถึง {fmt(high)} ใน {cols[idx[0]]['name']}" if iqr>0 else 'ช่วง IQR เป็นศูนย์ จึงไม่ระบุค่าผิดปกติด้วยวิธีนี้'
        elif kind=='distribution':
            vals=[number(r[idx[0]]) for r in rows if number(r[idx[0]]) is not None];bins=min(10,max(3,int(math.sqrt(len(vals)))));low=min(vals);step=(max(vals)-low)/bins;counts=[0]*bins
            if not step or not math.isfinite(step):raise ValueError('ช่วงตัวเลขไม่เหมาะกับการสร้าง histogram')
            for v in vals:counts[min(bins-1,int((v-low)/step))]+=1
            points=[{'label':f'{fmt(low+i*step)}–{fmt(low+(i+1)*step)}','value':n} for i,n in enumerate(counts)]
            data={'bins':points,'sample_size':len(vals),'method':'equal-width histogram; last bin includes upper bound'};chart={'type':'bar','points':points,'unit':'จำนวนค่า'};finding=f"แสดงการกระจายของ {cols[idx[0]]['name']} จาก {fmt(len(vals))} ค่า ใน {bins} ช่วง"
        elif kind in ('category','frequency','trend'):
            groups=defaultdict(list)
            for row in rows:
                key=iso(row[idx[0]]) if kind=='trend' else str(row[idx[0]]) if present(row[idx[0]]) else None
                if key is None:continue
                if kind=='trend':key=key[:7]
                value=1 if kind=='frequency' else number(row[idx[1]])
                if value is not None:groups[key].append(value)
            points=[{'label':k,'value':len(v) if kind=='frequency' else statistics.mean(v),'count':len(v)} for k,v in groups.items()]
            points=sorted(points,key=(lambda x:x['label']) if kind=='trend' else (lambda x:-x['value']))
            data={'groups':points,'aggregation':'count' if kind=='frequency' else 'mean','included':sum(x['count'] for x in points),'excluded':len(rows)-sum(x['count'] for x in points)}
            if kind=='trend':
                months=[int(x['label'][:4])*12+int(x['label'][5:7]) for x in points]
                data['missing_months']=max(months)-min(months)+1-len(months) if months else 0
            chart={'type':'line' if kind=='trend' else 'bar','points':points,'unit':'จำนวนแถว' if kind=='frequency' else f"ค่าเฉลี่ย {cols[idx[1]]['name']}"}
            if points:
                top=max(points,key=lambda x:x['value']);finding=f"{top['label']} มี{'จำนวนแถว' if kind=='frequency' else 'ค่าเฉลี่ย'}สูงสุด {fmt(top['value'])} (ข้อมูล {fmt(top['count'])} แถว)"
            else:finding='ไม่มีคู่ข้อมูลที่ใช้คำนวณได้'
        elif kind=='correlation':
            pairs=[(number(r[idx[0]]),number(r[idx[1]])) for r in rows if number(r[idx[0]]) is not None and number(r[idx[1]]) is not None]
            sx=max(abs(x) for x,y in pairs);sy=max(abs(y) for x,y in pairs)
            scaled=[(x/sx,y/sy) for x,y in pairs];xs=[x for x,y in scaled];ys=[y for x,y in scaled];mx=statistics.mean(xs);my=statistics.mean(ys)
            r=math.fsum((x-mx)*(y-my) for x,y in scaled)/(math.sqrt(math.fsum((x-mx)**2 for x in xs))*math.sqrt(math.fsum((y-my)**2 for y in ys)))
            if not math.isfinite(r):raise ValueError('ไม่สามารถคำนวณความสัมพันธ์อย่างเสถียรจากตัวเลขชุดนี้')
            data={'r':max(-1,min(1,r)),'sample_size':len(pairs),'excluded':len(rows)-len(pairs)};finding=f"Pearson r = {fmt(data['r'])} จาก {fmt(len(pairs))} คู่ข้อมูล ความสัมพันธ์นี้ไม่ยืนยันเหตุและผล"
            chart={'type':'scatter','points':[{'x':x,'y':y} for x,y in pairs[:1000]],'x_label':cols[idx[0]]['name'],'unit':cols[idx[1]]['name']}
        eid=f'EV-{len(evidence)+1:03d}'
        e={'evidence_id':eid,'calculation_id':sid,'type':kind,'source':{'table_id':table_id,'sheet':t['sheet'],'range':t['range'],'source_columns':p['source_columns']},'data':data,'finding':finding,'confidence':1.0,'method':p['reason'],'limitations':['ผลคำนวณขึ้นกับความถูกต้องของข้อมูลและบทบาทคอลัมน์ที่อนุมาน','ไม่ใช่การทดสอบนัยสำคัญทางสถิติ']}
        evidence.append(e);results.append({'id':sid,'title':p['title'],'type':kind,'finding':finding,'data':data,'chart':chart,'evidence_id':eid,'method':p['reason']})
    return {'metadata':{'title':'รายงานการวิเคราะห์ข้อมูล','table':t['name'],'source_range':t['range'],'objective':objective[:1000],'generated_at':datetime.now().isoformat(),'engine':'Python deterministic engine 1.0','interpretation_mode':'evidence-based templates'},'dataset_overview':{k:t[k] for k in ['rows_count','columns_count','sheet','range','columns']},'data_quality':t['quality'],'excluded_rows':t['excluded_rows'],'analyses':results,'evidence':evidence,'sections':[r['type'] for r in results],'executive_summary':[r['finding'] for r in results[:5]],'recommendations':['ตรวจสอบค่าว่าง แถวซ้ำ และบทบาทคอลัมน์ก่อนนำผลไปใช้ตัดสินใจ','ตรวจสอบเหตุผลของค่าผิดปกติกับเจ้าของข้อมูล โดยไม่ลบออกอัตโนมัติ'] if t['quality']['issues'] else ['ตรวจทานความหมายของตัวชี้วัดและหน่วยกับเจ้าของข้อมูลก่อนนำผลไปใช้'],'limitations':['ไม่มีการอนุมานเหตุและผล','ค่าเฉลี่ยอาจไม่เหมาะกับตัวชี้วัดทุกประเภท','วิเคราะห์เฉพาะตารางที่เลือก ไม่มีการ join ข้ามตาราง','วัตถุประสงค์บันทึกไว้ในรายงาน การจัดลำดับตามวัตถุประสงค์ต้องใช้ Gemini']}

def analyze_workbook(selected_types=None, objective='', progress=None):
    if not WORKBOOK or not TABLES:raise ValueError('ไม่พบชุดข้อมูล กรุณาอัปโหลดใหม่')
    allowed={p['type'] for p in WORKBOOK['opportunities']}
    selected=allowed if selected_types is None else set(selected_types)
    if selected-allowed:raise ValueError('ประเภทการวิเคราะห์ไม่ถูกต้อง')
    selected=selected|{'quality'}
    results=[];evidence=[];errors=[];table_reports=[]
    for index,t in enumerate(TABLES.values()):
        completed=0;failed=0
        plans=[p for p in t['opportunities'] if p['type'] in selected]
        for p in plans:
            try:
                part=analyze(t['id'],[p['id']],objective)
                eid=f'EV-{len(evidence)+1:03d}'
                ev={**part['evidence'][0],'evidence_id':eid}
                evidence.append(ev)
                results.append({**part['analyses'][0],'evidence_id':eid,'source':ev['source']})
                completed+=1
            except (ValueError,OverflowError,ZeroDivisionError,statistics.StatisticsError) as error:
                errors.append({'analysis_id':p['id'],'title':p['title'],'source':table_source(t),'message':str(error)})
                failed+=1
        table_reports.append({'id':t['id'],'name':t['name'],'sheet':t['sheet'],'range':t['range'],'rows_count':t['rows_count'],'columns_count':t['columns_count'],'analyses_count':completed,'errors_count':failed,'status':'partial' if failed else 'complete'})
        if progress:progress(index+1,len(TABLES),t['sheet'])
    summary=WORKBOOK['summary']
    sheets=[]
    for s in WORKBOOK['sheets']:
        members=[t for t in table_reports if t['sheet']==s['name']]
        failures=sum(t['errors_count'] for t in members)
        sheets.append({**s,'analyses_count':sum(t['analyses_count'] for t in members),'errors_count':failures,'status':('partial' if failures else 'complete') if members else 'no_table'})
    empty=sum(s['status']=='no_table' for s in sheets)
    overview=[f"อ่านครบ {len(sheets)} ชีต พบ {len(TABLES)} ตาราง รวม {fmt(summary['rows_count'])} แถวข้อมูล",f"คำนวณสำเร็จ {len(results)} รายการ จากทุกตารางที่ตรวจพบ"+(f" · คำนวณไม่สำเร็จ {len(errors)} รายการ กรุณาตรวจรายละเอียด" if errors else ''),f"ข้อมูลครบถ้วน {fmt(summary['quality']['completeness'])}% จาก {fmt(summary['cells_count'])} เซลล์ข้อมูล · พบค่าว่าง {fmt(summary['quality']['missing'])} เซลล์",'คำนวณแยกแต่ละตารางตามชนิดข้อมูล พร้อมระบุชีตและช่วงเซลล์ต้นทาง']
    if empty:overview.append(f'{empty} ชีตไม่มีตารางที่เข้าเกณฑ์ จึงแสดงเหตุผลไว้ในขอบเขตรายงาน')
    return {'metadata':{'title':'รายงานการวิเคราะห์ข้อมูลทุกชีต','table':f"ภาพรวม {len(sheets)} ชีต · {len(TABLES)} ตาราง",'source_range':'ทุกตารางที่ตรวจพบ','objective':objective[:1000],'generated_at':datetime.now().isoformat(),'engine':'Python deterministic engine 1.1','interpretation_mode':'evidence-based templates'},'dataset_overview':{'scope':'workbook','rows_count':summary['rows_count'],'columns_count':summary['columns_count'],'cells_count':summary['cells_count'],'sheets_count':len(sheets),'tables_count':len(TABLES),'sheet':'ทุกชีต','range':'ทุกตารางที่ตรวจพบ','columns':[],'sheets':sheets,'tables':table_reports},'data_quality':summary['quality'],'excluded_rows':[{**r,'source':table_source(t)} for t in TABLES.values() for r in t['excluded_rows']],'analyses':results,'evidence':evidence,'errors':errors,'sections':list(dict.fromkeys(r['type'] for r in results)),'executive_summary':overview,'recommendations':['ตรวจสอบประเด็นคุณภาพข้อมูลและตารางที่หาหัวตารางได้ด้วยความมั่นใจต่ำ','ตรวจทานความหมายและหน่วยของตัวชี้วัดกับเจ้าของข้อมูลก่อนใช้ตัดสินใจ'],'limitations':['คำนวณทุกตารางที่ตรวจพบโดยแยกกัน ไม่มีการ join หรือบวกตัวชี้วัดข้ามตาราง','ความครบถ้วนถ่วงตามจำนวนเซลล์ข้อมูล แถวซ้ำตรวจภายในแต่ละตารางและยังเก็บไว้','ชีตว่างหรือชีตที่ไม่มีตารางเข้าเกณฑ์แสดงในขอบเขต โดยไม่สร้างผลคำนวณแทน','ไม่คำนวณสูตร Excel ใหม่ และไม่อนุมานเหตุและผล','วัตถุประสงค์บันทึกในรายงานและใช้ประกอบคำตีความเมื่อเปิด Gemini']}

def dispatch(action, payload, progress=None):
    if action=='inspect':return inspect(bytes(payload['bytes']),payload['filename'])
    if action=='analyze':return analyze(payload['table_id'],payload['selected'],payload.get('objective',''))
    if action=='analyze_workbook':return analyze_workbook(payload.get('selected_types'),payload.get('objective',''),progress)
    raise ValueError('Unknown action')
