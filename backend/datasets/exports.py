"""Downloadable reports from persisted, deterministic dataset analysis.

CLI: exports.py export database.sqlite analysis.json pdf|xlsx|csv output [sheet_id]
Only the final JSON result/error is written to stdout. SQLite is opened read-only;
uploaded strings are always exported as inert text, never executable formulas.

PDF fonts: set PDF_FONT_PATH (and optionally PDF_FONT_BOLD_PATH) to a Thai-capable
TrueType font. Windows Leelawadee and Linux Noto Sans Thai are detected locally.
Linux packages commonly provide fonts-noto-core. Fonts are embedded, not bundled.
"""
from __future__ import annotations

import csv
import json
import math
import os
import re
import sqlite3
import sys
import unicodedata
from contextlib import closing
from pathlib import Path
from xml.sax.saxutils import escape


BRAND = "ASW Data Insight"
NAVY = "123F6D"
MIMES = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8"}


class ExportError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code, self.message = code, message


def text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return str(value)


def clean_text(value):
    # XML 1.0 (XLSX) excludes these controls. Keep newlines and tabs as text.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "\ufffd", text(value))


def formula_like(value):
    if not isinstance(value, str):
        return False
    probe = value
    while probe and (probe[0].isspace() or unicodedata.category(probe[0]) in ("Cc", "Cf")):
        probe = probe[1:]
    return bool(probe) and probe[0] in "=+-@"


def csv_value(value):
    if isinstance(value, str):
        return "'" + value if formula_like(value) else value
    return value


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def display(value):
    if number(value):
        return f"{value:,.4f}".rstrip("0").rstrip(".") if isinstance(value, float) else f"{value:,}"
    return clean_text(value)


def open_dataset(database):
    source = Path(database)
    if not source.is_file():
        raise ExportError("NOT_FOUND", "ไม่พบชุดข้อมูลสำหรับส่งออก กรุณาอัปโหลดและวิเคราะห์ใหม่")
    connection = sqlite3.connect(source.resolve().as_uri() + "?mode=ro", uri=True)
    connection.execute("PRAGMA query_only=ON")
    return connection


def dataset_metadata(connection):
    row = connection.execute("SELECT value FROM metadata WHERE key='dataset'").fetchone()
    if row is None:
        raise ExportError("INVALID_DATASET", "ชุดข้อมูลไม่มีข้อมูลต้นทางสำหรับส่งออก")
    dataset = json.loads(row[0])
    if not dataset.get("sheets"):
        raise ExportError("INVALID_DATASET", "ไม่พบชีตข้อมูลสำหรับส่งออก")
    return dataset


def selected_sheet(dataset, sheet_id):
    sheet_id = sheet_id or dataset["sheets"][0]["id"]
    if not isinstance(sheet_id, str) or not re.fullmatch(r"s\d+", sheet_id):
        raise ExportError("INVALID_SHEET", "ชีตที่เลือกสำหรับส่งออกไม่ถูกต้อง")
    sheet = next((item for item in dataset["sheets"] if item["id"] == sheet_id), None)
    if sheet is None:
        raise ExportError("INVALID_SHEET", "ไม่พบชีตที่เลือกสำหรับส่งออก")
    return sheet


def source_rows(connection, sheet):
    # Never interpolate a client-supplied identifier without checking metadata.
    identifier = sheet["id"]
    if not re.fullmatch(r"s\d+", identifier):
        raise ExportError("INVALID_SHEET", "รหัสชีตข้อมูลไม่ถูกต้อง")
    for (payload,) in connection.execute(f'SELECT data FROM "data_{identifier}" ORDER BY row_number'):
        record = json.loads(payload)
        yield [record.get(column["key"]) for column in sheet["columns"]]


def export_csv(connection, dataset, output, sheet_id=None):
    sheet = selected_sheet(dataset, sheet_id)
    # BOM lets Excel recognize Thai UTF-8 without an import dialog.
    with output.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow([csv_value(column["name"]) for column in sheet["columns"]])
        for row in source_rows(connection, sheet):
            writer.writerow([csv_value(value) for value in row])


def ai_summary(analysis):
    ai = analysis.get("ai") or {}
    if not isinstance(ai, dict):
        return ""
    narrative = ai.get("summary") or (ai.get("result") or {}).get("summary")
    return narrative if isinstance(narrative, str) else ""


def export_xlsx(connection, dataset, analysis, output):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.chart import BarChart, LineChart, ScatterChart, Series, Reference
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    workbook.properties.creator = BRAND
    workbook.properties.title = f"{BRAND} - {dataset['filename']}"
    navy = PatternFill("solid", fgColor=NAVY)
    pale = PatternFill("solid", fgColor="EFF4F9")

    def sheet(name, headers, widths):
        page = workbook.create_sheet(name)
        page.freeze_panes = "A2"
        page.sheet_properties.pageSetUpPr.fitToPage = True
        page.sheet_properties.outlinePr.summaryRight = False
        page.sheet_properties.tabColor = NAVY
        page.sheet_view.showGridLines = False
        for index, width in enumerate(widths, 1):
            page.column_dimensions[get_column_letter(index)].width = width
        row(page, headers, header=True)
        return page

    def row(page, values, header=False):
        cells = []
        for value in values:
            cell = WriteOnlyCell(page)
            if isinstance(value, str) or isinstance(value, (dict, list)):
                original = clean_text(value)
                if len(original) > 32767:
                    raise ExportError("XLSX_CELL_TOO_LONG", "มีข้อความยาวเกินขีดจำกัด 32,767 ตัวอักษรต่อเซลล์ของ Excel กรุณาส่งออกเป็น CSV")
                cell.value = original
                cell.data_type = "s"
                # Explicit string type already makes formulas inert; quotePrefix
                # also protects subsequent manual edits in Excel.
                cell.quotePrefix = formula_like(original)
            else:
                cell.value = value
            cell.font = Font(name="Leelawadee UI", size=11, color="FFFFFF" if header else "20344B", bold=header)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header:
                cell.fill = navy
            elif page.title.startswith("Data "):
                cell.alignment = Alignment(vertical="top")
            elif len(cells) == 0:
                cell.fill = pale
            if number(value):
                cell.number_format = "#,##0" if isinstance(value, int) else "#,##0.####"
            cells.append(cell)
        page.append(cells)

    overview = sheet("Overview", [BRAND, "Analysis report"], [34, 110])
    for label, value in [("Source file", dataset["filename"]), ("Generated at", analysis.get("generated_at")), ("Rows", dataset.get("rows_count")), ("Columns across sheets", dataset.get("columns_count")), ("Sheets", len(dataset["sheets"])), ("Summary", analysis.get("summary"))]:
        row(overview, [label, value])
    for section in (analysis.get("report") or {}).get("sections", []):
        row(overview, [section.get("title"), "\n\n".join(section.get("paragraphs", []))])
        if section.get("evidence_ids"):
            row(overview, ["Evidence", ", ".join(section["evidence_ids"])])
    row(overview, ["Methodology", "Statistics are computed from stored normalized records; source sheets are analyzed independently. Original formula cells remain text. Data tabs contain all stored records."])

    kpis = sheet("KPIs", ["Metric", "Value", "Display", "Source sheet", "Source column", "Method"], [34, 22, 24, 26, 28, 85])
    for item in analysis.get("kpis", []):
        source = item.get("source") or {}
        row(kpis, [item.get("name"), item.get("value"), item.get("formatted_value"), source.get("sheet"), source.get("column"), item.get("method")])

    insights = sheet("Insights", ["Evidence ID", "Priority", "Type", "Finding", "Description", "Metric", "Value", "Method", "Sheet", "Columns"], [18, 14, 20, 45, 85, 30, 24, 70, 24, 35])
    for item in analysis.get("insights", []):
        evidence = item.get("evidence") or {}
        row(insights, [item.get("id"), item.get("importance"), item.get("kind"), item.get("title"), item.get("description"), evidence.get("metric"), evidence.get("value"), evidence.get("method"), evidence.get("sheet"), ", ".join(evidence.get("columns", []))])

    quality = sheet("Quality", ["Sheet", "Rows", "Duplicate rows", "Missing cells", "Missing %", "Warnings"], [30, 18, 22, 22, 18, 100])
    columns = sheet("Columns", ["Sheet", "Column", "Type", "Missing", "Missing %", "Unique", "Count", "Minimum", "Maximum", "Mean", "Median", "Std deviation", "Sum", "Earliest", "Latest", "Potential outliers", "Top values"], [28, 35, 16, 16, 16, 16, 16, 20, 20, 20, 20, 20, 22, 28, 28, 20, 85])
    for profile in analysis.get("profiles", []):
        row(quality, [profile.get("sheet_name"), profile.get("rows_count"), profile.get("duplicate_rows"), profile.get("missing_count"), profile.get("missing_percentage"), "\n".join(profile.get("warnings", []))])
        for column in profile.get("columns", []):
            stats, dates = column.get("statistics") or {}, column.get("date_range") or {}
            row(columns, [profile.get("sheet_name"), column.get("name"), column.get("data_type"), column.get("missing_count"), column.get("missing_percentage"), column.get("unique_count"), stats.get("count"), stats.get("min"), stats.get("max"), stats.get("mean"), stats.get("median"), stats.get("std"), stats.get("sum"), dates.get("min"), dates.get("max"), (column.get("outliers") or {}).get("count"), column.get("top_values")])
    for warning in dataset.get("warnings", []):
        row(quality, ["Workbook", None, None, None, None, warning])

    charts = sheet("Charts", ["Chart", "X", "Y", "Method"], [48, 32, 24, 75])
    chart_row = 2
    for item in analysis.get("charts", []):
        points = item.get("data") or []
        if not points:
            continue
        first = chart_row
        for point in points:
            row(charts, [item.get("title"), point.get("x"), point.get("y"), item.get("method")])
            chart_row += 1
        if item.get("type") == "scatter" and all(number(point.get("x")) for point in points):
            chart = ScatterChart()
            chart.series.append(Series(Reference(charts, min_col=3, min_row=first, max_row=chart_row - 1), Reference(charts, min_col=2, min_row=first, max_row=chart_row - 1)))
        else:
            chart = LineChart() if item.get("type") == "line" else BarChart()
            chart.add_data(Reference(charts, min_col=3, min_row=first, max_row=chart_row - 1))
            chart.set_categories(Reference(charts, min_col=2, min_row=first, max_row=chart_row - 1))
        chart.title = item.get("title")
        chart.x_axis.title, chart.y_axis.title = item.get("x_label"), item.get("y_label")
        chart.legend = None
        chart.width, chart.height = 23, 10
        if chart.series:
            chart.series[0].graphicalProperties.solidFill = NAVY
            chart.series[0].graphicalProperties.line.solidFill = NAVY
        charts.add_chart(chart, f"F{first}")
        # Reserve enough vertical space for each native Excel chart.
        while chart_row < first + 22:
            row(charts, [])
            chart_row += 1

    for index, source in enumerate(dataset["sheets"], 1):
        name = re.sub(r"[\\/*?:\[\]]", "_", f"Data {index} {source['name']}")[:31].rstrip("'")
        page = sheet(name, [column["name"] for column in source["columns"]], [max(18, min(45, len(column["name"]) + 4)) for column in source["columns"]])
        for values in source_rows(connection, source):
            row(page, values)
    workbook.save(output)


def pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        import uharfbuzz  # noqa: F401 - required for reportlab Thai shaping
    except ImportError as exc:
        raise ExportError("MISSING_DEPENDENCY", "PDF ต้องใช้ uharfbuzz สำหรับภาษาไทย กรุณาติดตั้ง backend/requirements.txt") from exc
    configured = os.environ.get("PDF_FONT_PATH")
    candidates = [
        (configured, os.environ.get("PDF_FONT_BOLD_PATH") or configured),
        ("C:/Windows/Fonts/LeelawUI.ttf", "C:/Windows/Fonts/LeelaUIb.ttf"),
        ("C:/Windows/Fonts/leelawad.ttf", "C:/Windows/Fonts/leelawdb.ttf"),
        ("/usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf", "/usr/share/fonts/truetype/noto/NotoSansThai-Bold.ttf"),
        ("/usr/share/fonts/opentype/noto/NotoSansThai-Regular.ttf", "/usr/share/fonts/opentype/noto/NotoSansThai-Bold.ttf"),
        ("/usr/share/fonts/truetype/tlwg/Garuda.ttf", "/usr/share/fonts/truetype/tlwg/Garuda-Bold.ttf"),
    ]
    if configured and not Path(configured).is_file():
        raise ExportError("PDF_FONT_UNAVAILABLE", "ไม่พบไฟล์ฟอนต์ที่กำหนดใน PDF_FONT_PATH")
    for regular, bold in candidates:
        if regular and Path(regular).is_file():
            pdfmetrics.registerFont(TTFont("ReportThai", regular, shapable=True))
            pdfmetrics.registerFont(TTFont("ReportThaiBold", bold if bold and Path(bold).is_file() else regular, shapable=True))
            font = pdfmetrics.getFont("ReportThai")
            if not all(code in font.face.charToGlyph for code in (0x0E01, 0x0E34, 0x0E48)):
                raise ExportError("PDF_FONT_UNAVAILABLE", "ฟอนต์ PDF ไม่รองรับภาษาไทย กรุณากำหนด PDF_FONT_PATH เป็น Noto Sans Thai หรือ Leelawadee")
            pdfmetrics.registerFontFamily("ReportThai", normal="ReportThai", bold="ReportThaiBold", italic="ReportThai", boldItalic="ReportThaiBold")
            return "ReportThai", "ReportThaiBold"
    raise ExportError("PDF_FONT_UNAVAILABLE", "ไม่พบฟอนต์ภาษาไทย กรุณาติดตั้ง Noto Sans Thai หรือกำหนด PDF_FONT_PATH")


def export_pdf(dataset, analysis, output):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

    regular, bold = pdf_fonts()
    navy, muted = colors.HexColor("#" + NAVY), colors.HexColor("#52677E")
    styles = {
        "title": ParagraphStyle("Title", fontName=bold, fontSize=26, leading=35, textColor=navy, spaceAfter=16, shaping=1),
        "h1": ParagraphStyle("H1", fontName=bold, fontSize=15, leading=23, textColor=navy, spaceBefore=18, spaceAfter=8, keepWithNext=True, shaping=1),
        "body": ParagraphStyle("Body", fontName=regular, fontSize=10, leading=17, textColor=colors.HexColor("#22364B"), spaceAfter=8, shaping=1, splitLongWords=True),
        "small": ParagraphStyle("Small", fontName=regular, fontSize=8, leading=13, textColor=muted, spaceAfter=5, shaping=1),
        "cell": ParagraphStyle("Cell", fontName=regular, fontSize=8, leading=13, textColor=colors.HexColor("#22364B"), shaping=1),
        "header": ParagraphStyle("Header", fontName=bold, fontSize=8, leading=13, textColor=colors.white, shaping=1),
    }
    def paragraph(value, kind="body"):
        return Paragraph(escape(clean_text(value)).replace("\n", "<br/>"), styles[kind])

    def table(headers, rows, widths):
        contents = [[paragraph(value, "header") for value in headers]]
        contents.extend([paragraph(display(value), "cell") for value in row] for row in rows)
        result = Table(contents, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=1)
        result.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LINEBELOW", (0, -1), (-1, -1), .5, colors.HexColor("#DAE3EC")),
        ]))
        return result

    width, height = A4
    usable = width - 92
    story = [paragraph(BRAND, "title"), paragraph("รายงานวิเคราะห์ข้อมูล", "h1"), paragraph(dataset["filename"]), paragraph(f"วันที่วิเคราะห์: {analysis.get('generated_at', '')}", "small"), Spacer(1, 10)]
    story.append(table(["แถวข้อมูล", "คอลัมน์รวมทุกชีต", "ชีต"], [[dataset.get("rows_count"), dataset.get("columns_count"), len(dataset["sheets"])]], [usable / 3] * 3))
    story.extend([Spacer(1, 15), paragraph(analysis.get("summary", ""))])

    for section in (analysis.get("report") or {}).get("sections", []):
        story.append(paragraph(section.get("title", ""), "h1"))
        story.extend(paragraph(value) for value in section.get("paragraphs", []))
        if section.get("evidence_ids"):
            story.append(paragraph("หลักฐาน: " + ", ".join(section["evidence_ids"]), "small"))

    kpis = analysis.get("kpis", [])
    if kpis:
        story.extend([paragraph("ตัวชี้วัดและวิธีคำนวณ", "h1"), table(["ตัวชี้วัด / แหล่งข้อมูล", "ค่า", "วิธีคำนวณ"], [[f"{item.get('name', '')}\n{(item.get('source') or {}).get('sheet', '')}", item.get("formatted_value", item.get("value")), item.get("method")] for item in kpis], [usable * .35, usable * .2, usable * .45])])
    if analysis.get("insights"):
        story.append(paragraph("หลักฐานประกอบข้อค้นพบ", "h1"))
        for item in analysis["insights"]:
            evidence = item.get("evidence") or {}
            story.append(paragraph(f"{item.get('id', '')} | {item.get('title', '')}", "h1"))
            story.append(paragraph(item.get("description", "")))
            story.append(paragraph(f"{evidence.get('sheet', '')} | {evidence.get('metric', '')}: {display(evidence.get('value'))}\n{evidence.get('method', '')}", "small"))

    profiles = analysis.get("profiles", [])
    if profiles:
        story.extend([paragraph("คุณภาพข้อมูลแยกตามชีต", "h1"), table(["ชีต", "แถว", "แถวซ้ำ", "เซลล์ว่าง", "ว่าง (%)"], [[p.get("sheet_name"), p.get("rows_count"), p.get("duplicate_rows"), p.get("missing_count"), p.get("missing_percentage")] for p in profiles], [usable * .36] + [usable * .16] * 4)])
    story.append(paragraph("แหล่งข้อมูลและขอบเขตการวิเคราะห์", "h1"))
    story.append(paragraph(f"แหล่งข้อมูล: {dataset['filename']}"))
    story.append(paragraph("สถิติคำนวณจากทุกแถวข้อมูลที่จัดเก็บ โดยวิเคราะห์แต่ละชีตแยกกัน ไม่รวมข้อมูลต่างชีตโดยสมมติความสัมพันธ์เอง เซลล์สูตรต้นทางเก็บเป็นข้อความและไม่คำนวณซ้ำ ดาวน์โหลด Excel เพื่อดูรายละเอียดคอลัมน์ กราฟ และข้อมูลครบทุกแถว หรือ CSV สำหรับข้อมูลของชีตที่เลือก"))
    for source in dataset["sheets"]:
        story.append(paragraph(f"{source['name']}: {source['rows_count']:,} แถว / {len(source['columns'])} คอลัมน์", "small"))
        story.extend(paragraph(warning, "small") for warning in source.get("warnings", []))
    story.extend(paragraph(warning, "small") for warning in dataset.get("warnings", []))

    def page_frame(canvas, document):
        canvas.saveState()
        canvas.setFillColor(navy)
        canvas.rect(0, height - 9, width, 9, stroke=0, fill=1)
        canvas.setStrokeColor(colors.HexColor("#D9E2EC"))
        canvas.line(46, 40, width - 46, 40)
        canvas.setFont(regular, 8)
        canvas.setFillColor(muted)
        canvas.drawString(46, 26, BRAND)
        canvas.drawRightString(width - 46, 26, str(document.page))
        canvas.restoreState()

    document = SimpleDocTemplate(str(output), pagesize=A4, leftMargin=46, rightMargin=46, topMargin=42, bottomMargin=55, title=f"{BRAND} - {dataset['filename']}", author=BRAND, pageCompression=1)
    document.build(story, onFirstPage=page_frame, onLaterPages=page_frame)


def export_dashboard_pdf(payload_path, output_path):
    """Landscape dashboard PDF. Numbers come from the server-side query in the payload;
    chart pictures are optional JPEGs captured in the browser, never trusted for values."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    payload_path, output = Path(payload_path), Path(output_path)
    document = json.loads(payload_path.read_text(encoding="utf-8"))
    if not isinstance(document, dict) or not isinstance(document.get("charts"), list) or not isinstance(document.get("kpis"), list):
        raise ExportError("INVALID_REQUEST", "ข้อมูล Dashboard สำหรับ PDF ไม่ถูกต้อง")
    regular, bold = pdf_fonts()
    navy, muted, ink, line = colors.HexColor("#" + NAVY), colors.HexColor("#5C6F84"), colors.HexColor("#192D47"), colors.HexColor("#E3E9F0")
    style = lambda name, font, size, color=ink, **extra: ParagraphStyle(name, fontName=font, fontSize=size, leading=size * 1.55, textColor=color, shaping=1, splitLongWords=True, **extra)
    styles = {"title": style("t", bold, 20, navy, spaceAfter=4), "h": style("h", bold, 12, navy, spaceBefore=10, spaceAfter=6),
              "body": style("b", regular, 9.5), "small": style("s", regular, 8, muted), "kpi": style("k", bold, 17, navy), "cell": style("c", regular, 8)}
    text = lambda value, kind="body": Paragraph(escape(clean_text(value)).replace("\n", "<br/>"), styles[kind])
    page_width, page_height = landscape(A4)
    usable = page_width - 60

    generated = str(document.get("generated_at", ""))[:16].replace("T", " ")
    story = [text(document.get("title") or "Dashboard", "title")]
    if document.get("description"):
        story.append(text(document["description"]))
    story.append(text(f"ไฟล์ {document.get('filename', '')} · ชีต {document.get('sheet', '')} · สร้างเมื่อ {generated} UTC · ใช้ {display(document.get('rows_matched'))} จาก {display(document.get('rows_total'))} แถว", "small"))
    filters = document.get("filters") or []
    story.append(text("ตัวกรอง: " + ("; ".join(f"{item.get('label', '')}: {item.get('text', '')}" for item in filters) if filters else "ไม่มี (ข้อมูลทั้งหมด)"), "small"))
    story.append(Spacer(1, 8))

    kpis = document["kpis"][:6]
    if kpis:
        cells = [[text(item.get("label", ""), "small"), text(item.get("text", "—"), "kpi")] for item in kpis]
        table = Table([cells], colWidths=[usable / len(cells)] * len(cells))
        table.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .6, line), ("INNERGRID", (0, 0), (-1, -1), .6, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.extend([table, Spacer(1, 10)])

    def data_table(chart, limit):
        rows = list(chart.get("data") or [])[:limit]
        if chart.get("others"):
            rows.append({"x": chart["others"].get("label"), "y": chart["others"].get("y")})
        if not rows:
            return text("ไม่มีข้อมูลตามตัวกรองนี้", "small")
        head = [chart.get("x_name") or "กลุ่ม", "จำนวนแถว" if chart.get("type") == "histogram" else chart.get("y_name") or "จำนวนแถว"]
        table = Table([[text(value, "cell") for value in head]] + [[text(display(row.get("x")), "cell"), text(display(row.get("y")), "cell")] for row in rows], repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), .4, line), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF3F8"))]))
        return table

    cell_width = usable / 2 - 8
    panels = []
    for chart in document["charts"]:
        body = [text(chart.get("title", ""), "h")]
        image = chart.get("image")
        picture = Path(image) if isinstance(image, str) else None
        if picture and picture.resolve().parent == payload_path.resolve().parent and picture.is_file() and picture.read_bytes()[:3] == b"\xff\xd8\xff" and picture.stat().st_size <= 3 * 1024 * 1024:
            width, height = ImageReader(str(picture)).getSize()
            scale = min(cell_width / width, 210 / height)
            body.append(Image(str(picture), width=width * scale, height=height * scale))
        elif chart.get("error"):
            body.append(text(chart["error"], "small"))
        else:
            body.append(data_table(chart, 12))
        if chart.get("sampled"):
            body.append(text(f"กราฟแสดงตัวอย่างจุดจาก {display(chart.get('points_total'))} คู่ข้อมูล", "small"))
        panels.append(body)
    if panels:
        rows = [panels[index:index + 2] + ([[]] if len(panels[index:index + 2]) == 1 else []) for index in range(0, len(panels), 2)]
        grid = Table(rows, colWidths=[usable / 2] * 2)
        grid.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
        story.append(grid)

    insights = document.get("insights") or []
    if insights or document.get("ai_summary"):
        story.append(text("ข้อสังเกตสำคัญ", "h"))
        if document.get("ai_summary"):
            story.append(text(document["ai_summary"]))
        for item in insights:
            evidence = " | ".join(f"{e.get('id', '')} {e.get('method', '')}".strip() for e in item.get("evidence") or [])
            story.append(KeepTogether([text(f"• {item.get('title', '')}", "body"), text(item.get("description", "")), text(f"หลักฐาน: {evidence}", "small")]))

    story.append(text("ข้อมูลประกอบกราฟ", "h"))
    for chart in document["charts"]:
        story.append(KeepTogether([text(chart.get("title", ""), "body"), data_table(chart, 25), Spacer(1, 6)]))
    story.append(text("ตัวเลขทุกค่าคำนวณจากข้อมูลทุกแถวที่ตรงตามตัวกรอง ณ เวลาส่งออก ไม่นับแถวสรุปยอดซ้ำ", "small"))

    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(navy)
        canvas.rect(0, page_height - 7, page_width, 7, stroke=0, fill=1)
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(30, 18, BRAND)
        canvas.drawRightString(page_width - 30, 18, str(doc.page))
        canvas.restoreState()

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".part")
    try:
        SimpleDocTemplate(str(temporary), pagesize=(page_width, page_height), leftMargin=30, rightMargin=30, topMargin=26, bottomMargin=32,
                          title=clean_text(document.get("title") or "Dashboard"), author=BRAND, pageCompression=1).build(story, onFirstPage=frame, onLaterPages=frame)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return {"path": str(output.resolve()), "mime_type": "application/pdf", "filename": output.name}


def export_report(database, analysis_path, format_name, output_path, sheet_id=None):
    if format_name not in MIMES:
        raise ExportError("INVALID_FORMAT", "รองรับการส่งออก PDF, Excel และ CSV เท่านั้น")
    output = Path(output_path)
    if output.resolve() in {Path(database).resolve(), Path(analysis_path).resolve()}:
        raise ExportError("INVALID_REQUEST", "ตำแหน่งส่งออกต้องแยกจากข้อมูลต้นทาง")
    if not Path(analysis_path).is_file():
        raise ExportError("ANALYSIS_NOT_READY", "ผลวิเคราะห์ยังไม่พร้อมสำหรับส่งออก")
    analysis = json.loads(Path(analysis_path).read_text(encoding="utf-8-sig"))
    if format_name != "csv" and (not isinstance(analysis, dict) or not isinstance((analysis.get("report") or {}).get("sections"), list)):
        raise ExportError("ANALYSIS_NOT_READY", "ผลวิเคราะห์ยังไม่สมบูรณ์ กรุณาวิเคราะห์ใหม่")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".part")
    try:
        with closing(open_dataset(database)) as connection:
            dataset = dataset_metadata(connection)
            if format_name == "csv":
                export_csv(connection, dataset, temporary, sheet_id)
            elif format_name == "xlsx":
                export_xlsx(connection, dataset, analysis, temporary)
            else:
                export_pdf(dataset, analysis, temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return {"path": str(output.resolve()), "mime_type": MIMES[format_name], "filename": output.name}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        if len(sys.argv) == 4 and sys.argv[1] == "dashboard-pdf":
            result = export_dashboard_pdf(sys.argv[2], sys.argv[3])
        elif len(sys.argv) not in (6, 7) or sys.argv[1] != "export":
            raise ExportError("INVALID_REQUEST", "คำสั่งส่งออกไม่ถูกต้อง")
        else:
            result = export_report(*sys.argv[2:])
        print(json.dumps({"result": result}, ensure_ascii=False, allow_nan=False))
    except ExportError as exc:
        print(json.dumps({"error": {"code": exc.code, "message": exc.message}}, ensure_ascii=False))
        return 1
    except ImportError:
        print(json.dumps({"error": {"code": "MISSING_DEPENDENCY", "message": "กรุณาติดตั้ง backend/requirements.txt ก่อนส่งออกรายงาน"}}, ensure_ascii=False))
        return 1
    except Exception:
        # Do not expose paths, source rows or backend exceptions to the browser.
        print(json.dumps({"error": {"code": "EXPORT_FAILED", "message": "สร้างไฟล์ส่งออกไม่สำเร็จ กรุณาลองอีกครั้งหรือตรวจสอบข้อมูลและการตั้งค่าเซิร์ฟเวอร์"}}, ensure_ascii=False))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
