"""Bounded CSV/XLSX ingestion and paginated SQLite preview.

This worker never evaluates cells or calls an LLM. Its stdout is a JSON-lines
protocol; uploaded values are returned only by the explicit preview command.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import sqlite3
import sys
import zipfile
from xml.etree import ElementTree
from datetime import date, datetime, timezone
from pathlib import Path


DEFAULT_LIMITS = {
    "max_rows": 100_000,
    "max_columns": 200,
    "max_cells": 2_000_000,
    "max_uncompressed_bytes": 100_000_000,
    "max_sheets": 50,
}
MAX_SAFE_INTEGER = 9_007_199_254_740_991
NUMBER = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z")
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}(?:T\d{2}:\d{2}.*)?\Z")
# A summary line starts with a total word ("รวม (Total)", "Grand Total ...", "VAT 7%").
TOTAL_ROW = re.compile(r"(?:grand\s*total|sub\s*-?\s*total|totals?\b|รวม|ยอดรวม|vat\b|ภาษีมูลค่าเพิ่ม)", re.I)
EXCEL_ERROR = re.compile(r"#(?:DIV/0!|N/A|NAME\?|NULL!|NUM!|REF!|VALUE!|SPILL!|CALC!)")


class DatasetError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code
        self.message = message


def fail(code, message):
    raise DatasetError(code, message)


def emit(value):
    print(json.dumps(value, ensure_ascii=False, allow_nan=False), flush=True)


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def limits_from(values):
    if not isinstance(values, dict):
        fail("INVALID_REQUEST", "การตั้งค่าขีดจำกัดต้องเป็น JSON object")
    result = dict(DEFAULT_LIMITS)
    for key in result:
        value = values.get(key, result[key])
        if isinstance(value, bool) or not isinstance(value, int) or value < 1:
            fail("INVALID_REQUEST", f"ขีดจำกัด {key} ต้องเป็นจำนวนเต็มบวก")
        result[key] = value
    return result


def is_iso_date(value):
    if not ISO_DATE.fullmatch(value):
        return False
    try:
        if "T" in value:
            datetime.fromisoformat(value.replace("Z", "+00:00"))
        else:
            date.fromisoformat(value)
        return True
    except ValueError:
        return False


def normalize(value, context):
    """Keep IDs, ambiguous dates, formula text and original text as text."""
    if blank(value):
        return None, "empty"
    if isinstance(value, bool):
        return value, "boolean"
    if isinstance(value, (date, datetime)):
        return value.isoformat(), "date"
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            fail("UNSUPPORTED_VALUE", f"{context}: พบตัวเลขที่ไม่เป็นค่าจำกัด กรุณาแก้ไขเซลล์นี้")
        if isinstance(value, int) and abs(value) > MAX_SAFE_INTEGER:
            return str(value), "text"
        return value, "number"
    if not isinstance(value, str):
        fail("UNSUPPORTED_VALUE", f"{context}: พบชนิดข้อมูลที่ยังไม่รองรับ กรุณาแปลงเป็นข้อความหรือตัวเลข")
    if any(ord(char) < 32 and char not in "\r\n\t" for char in value):
        fail("UNSUPPORTED_VALUE", f"{context}: พบอักขระไบนารีหรืออักขระควบคุมที่ไม่รองรับ")
    trimmed = value.strip()
    if trimmed.casefold() in ("true", "false"):
        return trimmed.casefold() == "true", "boolean"
    if is_iso_date(trimmed):
        return trimmed, "date"
    if NUMBER.fullmatch(trimmed) and not re.match(r"[+-]?0\d", trimmed):
        try:
            number = float(trimmed) if any(c in trimmed for c in ".eE") else int(trimmed)
            if math.isfinite(number) and abs(number) <= MAX_SAFE_INTEGER:
                return number, "number"
        except (ValueError, OverflowError):
            pass
    return value, "text"


def validate_headers(values, sheet, row_number, max_columns):
    if len(values) > max_columns:
        fail("LIMIT_EXCEEDED", f"ชีต {sheet}: มีคอลัมน์เกินขีดจำกัด {max_columns:,} คอลัมน์ กรุณาแบ่งไฟล์")
    if not values or any(blank(v) or not isinstance(v, str) for v in values):
        fail("MISSING_HEADERS", f"ชีต {sheet} แถว {row_number}: หัวคอลัมน์ขาดหายหรือไม่ใช่ข้อความ กรุณาใส่ชื่อให้ครบทุกคอลัมน์ในแถวแรกที่มีข้อมูล")
    names = [value.strip() for value in values]
    if all(NUMBER.fullmatch(name) or is_iso_date(name) for name in names):
        fail("MISSING_HEADERS", f"ชีต {sheet} แถว {row_number}: ไม่พบชื่อหัวคอลัมน์ กรุณาเพิ่มแถวชื่อคอลัมน์ก่อนข้อมูล")
    seen = set()
    for name in names:
        if any(ord(char) < 32 for char in name) or len(name) > 500:
            fail("INVALID_HEADERS", f"ชีต {sheet} แถว {row_number}: ชื่อคอลัมน์มีอักขระที่ไม่รองรับหรือยาวเกิน 500 ตัวอักษร")
        if name.casefold() in seen:
            fail("DUPLICATE_HEADERS", f"ชีต {sheet} แถว {row_number}: มีชื่อคอลัมน์ซ้ำ กรุณาตั้งชื่อแต่ละคอลัมน์ให้แตกต่างกัน")
        seen.add(name.casefold())
    return names


def csv_rows(path):
    with path.open("rb") as raw:
        prefix = raw.read(4)
        raw.seek(0)
        encoding = "utf-16" if prefix.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        with io.TextIOWrapper(raw, encoding=encoding, errors="strict", newline="") as stream:
            sample = stream.read(65_536)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                first = next((line for line in sample.splitlines() if line.strip()), "")
                delimiter = max(",;\t|", key=first.count)
                dialect = type("DetectedCSV", (csv.excel,), {"delimiter": delimiter})
            reader = csv.reader(stream, dialect, strict=True)
            while True:
                row_number = reader.line_num + 1
                try:
                    values = next(reader)
                except StopIteration:
                    break
                yield row_number, values, 0


def xlsx_preflight(path, limits):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 10_000 or len(set(names)) != len(names):
                fail("INVALID_FILE", "โครงสร้างไฟล์ XLSX ไม่ถูกต้องหรือซับซ้อนเกินขีดจำกัด")
            if sum(entry.file_size for entry in entries) > limits["max_uncompressed_bytes"]:
                fail("LIMIT_EXCEEDED", f"ข้อมูล XLSX หลังคลายไฟล์เกิน {limits['max_uncompressed_bytes'] / 1_000_000:g} MB กรุณาแบ่งไฟล์")
            if any(entry.flag_bits & 1 for entry in entries):
                fail("INVALID_FILE", "ไม่รองรับไฟล์ XLSX ที่เข้ารหัส กรุณาบันทึกสำเนาที่ไม่ใส่รหัสผ่าน")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                fail("INVALID_FILE", "ไฟล์นี้ไม่ใช่สมุดงาน XLSX ที่ถูกต้อง กรุณาบันทึกใหม่จาก Excel")
            if any("vba" in name.casefold() or "macrosheet" in name.casefold() for name in names):
                fail("UNSUPPORTED_FORMAT", "ไม่รองรับสมุดงานที่มี macro กรุณาบันทึกเป็น XLSX ที่ไม่มี macro")
            content_types = archive.read("[Content_Types].xml")
            if b"macroenabled" in content_types.lower() or b"macrosheet" in content_types.lower():
                fail("UNSUPPORTED_FORMAT", "ไม่รองรับสมุดงานที่มี macro กรุณาบันทึกเป็น XLSX ที่ไม่มี macro")
            if archive.testzip() is not None:
                fail("INVALID_FILE", "ไฟล์ XLSX เสียหาย กรุณาเปิดไฟล์ใน Excel แล้วบันทึกใหม่")
            # read_only iteration can silently skip repeated/out-of-order rows
            # and synthesize enormous gaps for forged coordinates. Validate
            # worksheet structure before trusting that iterator.
            for name in names:
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    with archive.open(name) as source:
                        validate_sheet_coordinates(source)
    except (zipfile.BadZipFile, OSError, RuntimeError, KeyError, ElementTree.ParseError):
        fail("INVALID_FILE", "อ่านโครงสร้างไฟล์ XLSX ไม่สำเร็จ กรุณาตรวจว่าไฟล์ไม่เสียหาย")


def validate_sheet_coordinates(source):
    previous_row = 0
    current_row = None
    previous_column = 0
    for event, element in ElementTree.iterparse(source, events=("start", "end")):
        tag = element.tag.rsplit("}", 1)[-1]
        if event == "start" and tag == "row":
            row = element.get("r", str(previous_row + 1))
            if not row.isascii() or not row.isdigit() or not previous_row < int(row) <= 1_048_576:
                fail("INVALID_FILE", "โครงสร้าง XLSX มีเลขแถวซ้ำ ผิดลำดับ หรือเกินขอบเขต Excel กรุณาเปิดไฟล์แล้วบันทึกใหม่")
            current_row = previous_row = int(row)
            previous_column = 0
        elif event == "start" and tag == "c":
            if current_row is None:
                fail("INVALID_FILE", "โครงสร้างเซลล์ XLSX ไม่ถูกต้อง กรุณาเปิดไฟล์ใน Excel แล้วบันทึกใหม่")
            coordinate = element.get("r")
            column = previous_column + 1
            if coordinate:
                match = re.fullmatch(r"([A-Za-z]{1,3})([1-9][0-9]{0,6})", coordinate)
                if not match or int(match[2]) != current_row:
                    fail("INVALID_FILE", "ตำแหน่งเซลล์ XLSX ไม่ตรงกับแถว กรุณาเปิดไฟล์ใน Excel แล้วบันทึกใหม่")
                column = 0
                for char in match[1].upper():
                    column = column * 26 + ord(char) - ord("A") + 1
            if not previous_column < column <= 16_384:
                fail("INVALID_FILE", "โครงสร้าง XLSX มีคอลัมน์ซ้ำ ผิดลำดับ หรือเกินขอบเขต Excel กรุณาเปิดไฟล์แล้วบันทึกใหม่")
            previous_column = column
        elif event == "end":
            if tag == "row":
                current_row = None
            element.clear()


def xlsx_rows(sheet, formula_sheet, uncached):
    """Yield the values Excel last calculated. Formulas are never evaluated here;
    a formula cell without a saved result is stored empty and counted in `uncached`."""
    import itertools
    # Ignore inaccurate worksheet dimensions, including styled empty tails.
    sheet.reset_dimensions()
    formula_sheet.reset_dimensions()
    for index, (cells, formula_cells) in enumerate(itertools.zip_longest(sheet.iter_rows(), formula_sheet.iter_rows(), fillvalue=()), 1):
        values = []
        formulas = 0
        for cell, source in itertools.zip_longest(cells, formula_cells):
            value = cell.value if cell is not None else None
            if cell is not None and cell.data_type == "e" or isinstance(value, str) and EXCEL_ERROR.fullmatch(value):
                value = None
                uncached["errors"] = uncached.get("errors", 0) + 1
            if source is not None and source.data_type == "f":
                formulas += 1
                if value is None:
                    uncached[sheet.title] = uncached.get(sheet.title, 0) + 1
            values.append(value)
        while values and blank(values[-1]):
            values.pop()
        yield index, values, formulas


def xls_rows(book, sheet):
    import xlrd
    for index in range(sheet.nrows):
        values = []
        for cell in sheet.row(index):
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                values.append(value.date() if value.time() == datetime.min.time() else value)
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                values.append(int(cell.value) if float(cell.value).is_integer() and abs(cell.value) <= MAX_SAFE_INTEGER else cell.value)
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                values.append(bool(cell.value))
            elif cell.ctype == xlrd.XL_CELL_TEXT:
                values.append(cell.value)
            else:
                values.append(None)
        while values and blank(values[-1]):
            values.pop()
        yield index + 1, values, 0


def resolve_sheet_header(iterator, sheet_name, limits):
    buffer = []
    blank_rows = 0
    for row_number, values, formulas in iterator:
        while values and blank(values[-1]):
            values.pop()
        if not any(not blank(value) for value in values):
            blank_rows += 1
            continue
        buffer.append((row_number, values, formulas))
        if len(buffer) >= 30:
            break

    if not buffer:
        return None, [], None, blank_rows, []

    max_width = max(len(v) for _, v, _ in buffer)
    first_row_num, first_values, _ = buffer[0]

    # Check if the first row is a title/banner row rather than the table header
    # (e.g. 1 cell title when table has 3+ columns, or 1-2 cell banner when table has 4+ columns).
    is_title_banner = max_width >= 3 and len(first_values) < max_width and len(first_values) <= (1 if max_width == 3 else 2)

    chosen_idx = 0
    chosen_names = None

    if is_title_banner:
        for idx in range(1, len(buffer)):
            row_number, values, _ = buffer[idx]
            if len(values) >= max(3, int(max_width * 0.7)):
                try:
                    chosen_names = validate_headers(values, sheet_name, row_number, limits["max_columns"])
                    chosen_idx = idx
                    break
                except DatasetError:
                    pass

    if chosen_names is None:
        chosen_names = validate_headers(first_values, sheet_name, first_row_num, limits["max_columns"])
        chosen_idx = 0

    header_row_num, _, _ = buffer[chosen_idx]
    skipped_rows = buffer[:chosen_idx]
    remaining_rows = buffer[chosen_idx + 1:]

    warnings = []
    if skipped_rows:
        warnings.append(f"ข้ามข้อความส่วนหัว {len(skipped_rows)} แถวก่อนเริ่มตารางข้อมูล (ใช้แถว {header_row_num} เป็นหัวตาราง)")

    return chosen_names, remaining_rows, header_row_num, blank_rows, warnings


def ingest(input_path, sqlite_path, filename, limits_values=None, progress=None):
    import itertools
    limits = limits_from({} if limits_values is None else limits_values)
    input_path, sqlite_path = Path(input_path), Path(sqlite_path)
    progress = progress or (lambda stage, value: None)
    progress("validating", 20)
    extension = Path(filename).suffix.casefold()
    if extension not in (".csv", ".xlsx", ".xls"):
        fail("UNSUPPORTED_FORMAT", "รองรับไฟล์ CSV, XLSX และ XLS เท่านั้น")
    if not input_path.is_file() or not input_path.stat().st_size:
        fail("EMPTY_FILE", "ไฟล์ว่าง กรุณาเลือกไฟล์ที่มีหัวคอลัมน์และแถวข้อมูล")
    if sqlite_path.exists():
        fail("INVALID_REQUEST", "พื้นที่เก็บข้อมูลชุดนี้มีอยู่แล้ว กรุณาเริ่มการอัปโหลดใหม่")
    book = None
    formula_book = None
    connection = None
    succeeded = False
    uncached = {}
    try:
        if extension == ".xlsx":
            xlsx_preflight(input_path, limits)
            import openpyxl
            book = openpyxl.load_workbook(input_path, read_only=True, data_only=True, keep_links=False)
            formula_book = openpyxl.load_workbook(input_path, read_only=True, data_only=False, keep_links=False)
            if len(book.worksheets) > limits["max_sheets"]:
                fail("LIMIT_EXCEEDED", f"สมุดงานมีชีตเกินขีดจำกัด {limits['max_sheets']} ชีต กรุณาแบ่งไฟล์")
            sources = [(sheet.title, xlsx_rows(sheet, formulas, uncached), sheet.sheet_state) for sheet, formulas in zip(book.worksheets, formula_book.worksheets)]
        elif extension == ".xls":
            import xlrd
            book = xlrd.open_workbook(str(input_path), on_demand=True)
            if book.nsheets > limits["max_sheets"]:
                fail("LIMIT_EXCEEDED", f"สมุดงานมีชีตเกินขีดจำกัด {limits['max_sheets']} ชีต กรุณาแบ่งไฟล์")
            sheets = [book.sheet_by_index(index) for index in range(book.nsheets)]
            if sum(sheet.nrows * sheet.ncols for sheet in sheets) > limits["max_cells"] * 2:
                fail("LIMIT_EXCEEDED", f"จำนวนเซลล์ข้อมูลรวมเกิน {limits['max_cells']:,} เซลล์ กรุณาแบ่งไฟล์")
            sources = [(sheet.name, xls_rows(book, sheet), "visible" if sheet.visibility == 0 else "hidden") for sheet in sheets]
        else:
            sources = [("CSV", csv_rows(input_path), "visible")]
        progress("reading", 35)
        connection = sqlite3.connect(sqlite_path)
        connection.execute("CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        result = {"filename": filename, "rows_count": 0, "columns_count": 0, "sheets": [], "warnings": [], "created_at": datetime.now(timezone.utc).isoformat()}
        total_cells = 0
        for source_index, (sheet_name, iterator, visibility) in enumerate(sources):
            progress("understanding_columns", 50)
            names, remaining_buffer, header_row, blank_rows, header_warnings = resolve_sheet_header(iterator, sheet_name, limits)
            if names is None:
                result["warnings"].append(f"ชีต {sheet_name}: ข้ามชีตว่างเพราะไม่มีข้อมูล")
                continue

            columns = [{"key": f"c{i}", "name": name, "data_type": "empty"} for i, name in enumerate(names)]
            kinds = [set() for _ in names]
            rows_count = 0
            formulas_count = 0
            total_rows, total_labels = [], []
            sheet_id = f"s{len(result['sheets'])}"
            connection.execute(f'CREATE TABLE "data_{sheet_id}" (row_number INTEGER PRIMARY KEY, data TEXT NOT NULL)')

            for row_number, values, formulas in itertools.chain(remaining_buffer, iterator):
                while values and blank(values[-1]):
                    values.pop()
                if not any(not blank(value) for value in values):
                    blank_rows += 1
                    continue
                if len(values) > len(names):
                    if all(blank(v) for v in values[len(names):]):
                        values = values[:len(names)]
                    else:
                        fail("MISSING_HEADERS", f"ชีต {sheet_name} แถว {row_number}: มีข้อมูลเกินจำนวนหัวคอลัมน์ กรุณาเติมหัวคอลัมน์หรือแก้จำนวนช่องให้ตรงกัน")
                if result["rows_count"] + rows_count + 1 > limits["max_rows"]:
                    fail("LIMIT_EXCEEDED", f"จำนวนแถวข้อมูลรวมเกิน {limits['max_rows']:,} แถว กรุณาแบ่งไฟล์")
                total_cells += len(names)
                if total_cells > limits["max_cells"]:
                    fail("LIMIT_EXCEEDED", f"จำนวนเซลล์ข้อมูลรวมเกิน {limits['max_cells']:,} เซลล์ กรุณาแบ่งไฟล์")
                first = next((value for value in values if not blank(value)), None)
                if isinstance(first, str) and TOTAL_ROW.match(first.strip()):
                    total_rows.append(row_number)
                    total_labels.append(first.strip()[:60])
                record = {}
                for index in range(len(names)):
                    value, kind = normalize(values[index] if index < len(values) else None, f"ชีต {sheet_name} แถว {row_number} คอลัมน์ {index + 1}")
                    record[f"c{index}"] = value
                    if kind != "empty":
                        kinds[index].add(kind)
                connection.execute(f'INSERT INTO "data_{sheet_id}" VALUES (?, ?)', (row_number, json.dumps(record, ensure_ascii=False, allow_nan=False)))
                rows_count += 1
                formulas_count += formulas
                if rows_count % 5000 == 0:
                    progress("reading", min(70, 35 + int((source_index + .5) / len(sources) * 35)))

            if not rows_count:
                fail("EMPTY_DATASET", f"ชีต {sheet_name}: พบเฉพาะหัวคอลัมน์ กรุณาเพิ่มแถวข้อมูลอย่างน้อยหนึ่งแถว")
            warnings = list(header_warnings)
            if blank_rows:
                warnings.append(f"ข้ามแถวว่าง {blank_rows:,} แถว โดยคงเลขแถวต้นฉบับไว้")
            if formulas_count:
                warnings.append(f"ใช้ค่าที่ Excel คำนวณไว้ล่าสุดของสูตร {formulas_count:,} เซลล์ ระบบไม่คำนวณสูตรใหม่")
            if uncached.get(sheet_name):
                warnings.append(f"สูตร {uncached[sheet_name]:,} เซลล์ไม่มีค่าที่คำนวณไว้ จึงเก็บเป็นค่าว่าง กรุณาเปิดไฟล์ใน Excel แล้วบันทึกใหม่ก่อนอัปโหลด")
            if total_rows:
                warnings.append(f"ไม่นำแถวสรุปยอด {len(total_rows):,} แถวมาคำนวณ ({', '.join(total_labels[:4])}) เพื่อไม่ให้ยอดซ้ำ แถวเหล่านี้ยังแสดงในตารางข้อมูล")
            if visibility != "visible":
                warnings.append("ชีตนี้ถูกซ่อนในไฟล์ต้นฉบับและรวมอยู่ในข้อมูลที่อ่านแล้ว")
            for column, types in zip(columns, kinds):
                column["data_type"] = next(iter(types)) if len(types) == 1 else "mixed" if types else "empty"
            result["sheets"].append({"id": sheet_id, "name": sheet_name, "rows_count": rows_count, "columns": columns, "header_row": header_row, "warnings": warnings, "summary_rows": total_rows[:1000]})
            result["rows_count"] += rows_count
            result["columns_count"] += len(columns)
            progress("reading", 35 + int((source_index + 1) / len(sources) * 35))
        if uncached.get("errors"):
            result["warnings"].append(f"พบเซลล์ข้อผิดพลาดของสูตร Excel (เช่น #DIV/0!) {uncached['errors']:,} เซลล์ เก็บเป็นค่าว่างและไม่นำมาคำนวณ")
        if not result["sheets"]:
            fail("EMPTY_DATASET", "ไม่พบแถวข้อมูล กรุณาเลือกไฟล์ที่มีหัวคอลัมน์และข้อมูลอย่างน้อยหนึ่งแถว")
        combine_sheets(connection, result)
        progress("detecting_types", 80)
        connection.execute("INSERT INTO metadata VALUES ('dataset', ?)", (json.dumps(result, ensure_ascii=False, allow_nan=False),))
        connection.commit()
        progress("preview", 95)
        succeeded = True
        return result
    except DatasetError:
        raise
    except UnicodeError:
        fail("INVALID_ENCODING", "อ่านรหัสอักขระ CSV ไม่สำเร็จ กรุณาบันทึกเป็น CSV UTF-8 หรือ UTF-16 ที่มี BOM")
    except csv.Error:
        fail("INVALID_FILE", "โครงสร้าง CSV ไม่ถูกต้อง กรุณาตรวจเครื่องหมายคำพูด ตัวคั่น และขนาดข้อความในแต่ละเซลล์")
    except ImportError:
        fail("PARSER_UNAVAILABLE", "เซิร์ฟเวอร์ยังไม่มีเครื่องมืออ่าน XLSX กรุณาให้ผู้ดูแลติดตั้ง Python dependencies ของ backend แล้วลองใหม่")
    except (OSError, sqlite3.Error):
        fail("PROCESSING_FAILED", "อ่านหรือจัดเก็บข้อมูลไม่สำเร็จ กรุณาลองอัปโหลดใหม่")
    except Exception:
        fail("INVALID_FILE", "ไม่สามารถอ่านสมุดงานได้ กรุณาตรวจโครงสร้างหรือเปิดไฟล์ใน Excel แล้วบันทึกใหม่")
    finally:
        for workbook in (book, formula_book):
            if workbook is not None:
                (workbook.release_resources if hasattr(workbook, "release_resources") else workbook.close)()
        if connection is not None:
            connection.close()
        if not succeeded and connection is not None:
            sqlite_path.unlink(missing_ok=True)


def combine_sheets(connection, result):
    """Stack sheets that share the same header row into one extra table so a
    dashboard can compare them (e.g. one BOQ sheet per building). Adds a "ชีต"
    column and, when names follow PREFIX_SUFFIX, a "กลุ่มชีต" column. Summary
    rows stay out; source sheets are untouched and dataset totals ignore it."""
    groups = {}
    for sheet in result["sheets"]:
        signature = tuple(column["name"].casefold() for column in sheet["columns"])
        if len(signature) >= 2:
            groups.setdefault(signature, []).append(sheet)
    members = max((group for group in groups.values() if len(group) >= 2), key=lambda group: (sum(s["rows_count"] for s in group), len(group)), default=None)
    if not members:
        return
    width = len(members[0]["columns"])
    names = [sheet["name"] for sheet in members]
    prefixes = [name.split("_", 1)[0] for name in names]
    by_prefix = all("_" in name for name in names) and 2 <= len(set(prefixes)) < len(names)
    sheet_id = f"s{len(result['sheets'])}"
    columns = [{**column} for column in members[0]["columns"]] + [{"key": f"c{width}", "name": "ชีต", "data_type": "text"}]
    if by_prefix:
        columns.append({"key": f"c{width + 1}", "name": "กลุ่มชีต", "data_type": "text"})
    connection.execute(f'CREATE TABLE "data_{sheet_id}" (row_number INTEGER PRIMARY KEY, data TEXT NOT NULL)')
    rows = 0
    for sheet, prefix in zip(members, prefixes):
        # Identical header rows give identical positional keys (c0, c1, ...).
        extra = f"json_set(data, '$.c{width}', ?" + (f", '$.c{width + 1}', ?)" if by_prefix else ")")
        params = [sheet["name"], prefix] if by_prefix else [sheet["name"]]
        cursor = connection.execute(f'INSERT INTO "data_{sheet_id}" (data) SELECT {extra} FROM "data_{sheet["id"]}" WHERE row_number NOT IN (SELECT value FROM json_each(?)) ORDER BY row_number', [*params, json.dumps(sheet.get("summary_rows", []))])
        rows += cursor.rowcount
    for column in columns:
        types = {sheet["columns"][index]["data_type"] for sheet in members for index, candidate in enumerate(sheet["columns"]) if candidate["key"] == column["key"]}
        if types:
            column["data_type"] = next(iter(types)) if len(types) == 1 else "mixed"
    result["sheets"].append({"id": sheet_id, "name": f"รวมทุกชีต ({len(members)} ชีต)", "rows_count": rows, "columns": columns, "header_row": None,
                             "warnings": [f"รวมข้อมูลจาก {len(members)} ชีตที่มีหัวคอลัมน์เหมือนกัน: {', '.join(names[:8])}{' …' if len(names) > 8 else ''} โดยไม่นับแถวสรุปยอด"],
                             "summary_rows": [], "combined_from": names})


def preview(sqlite_path, query):
    if not isinstance(query, dict):
        fail("INVALID_REQUEST", "พารามิเตอร์ค้นหาต้องเป็น JSON object")
    if not Path(sqlite_path).is_file():
        fail("NOT_FOUND", "ไม่พบชุดข้อมูลนี้ กรุณาอัปโหลดใหม่")
    page, page_size = query.get("page", 1), query.get("page_size", 25)
    if isinstance(page, bool) or not isinstance(page, int) or page < 1 or page > 1_000_000:
        fail("INVALID_REQUEST", "หมายเลขหน้าต้องเป็นจำนวนเต็มตั้งแต่ 1 ถึง 1,000,000")
    if isinstance(page_size, bool) or not isinstance(page_size, int) or not 1 <= page_size <= 100:
        fail("INVALID_REQUEST", "จำนวนแถวต่อหน้าต้องอยู่ระหว่าง 1 ถึง 100")
    search = query.get("search", "")
    if not isinstance(search, str) or len(search) > 500:
        fail("INVALID_REQUEST", "คำค้นหาต้องเป็นข้อความไม่เกิน 500 ตัวอักษร")
    direction = query.get("direction", "asc")
    if direction not in ("asc", "desc"):
        fail("INVALID_REQUEST", "ทิศทางเรียงลำดับต้องเป็น asc หรือ desc")
    connection = sqlite3.connect(Path(sqlite_path).resolve().as_uri() + "?mode=ro", uri=True)
    try:
        connection.create_function("casefold", 1, lambda value: str(value).casefold() if value is not None else "", deterministic=True)
        metadata = connection.execute("SELECT value FROM metadata WHERE key='dataset'").fetchone()
        if not metadata:
            fail("NOT_FOUND", "ไม่พบข้อมูลชุดนี้ กรุณาอัปโหลดใหม่")
        dataset = json.loads(metadata[0])
        sheet_id = query.get("sheet", dataset["sheets"][0]["id"])
        sheet = next((item for item in dataset["sheets"] if item["id"] == sheet_id), None)
        if sheet is None or not re.fullmatch(r"s\d+", sheet_id):
            fail("INVALID_REQUEST", "ไม่พบชีตที่ระบุ กรุณาเลือกชีตจากรายการ")
        columns = {column["key"]: column for column in sheet["columns"]}
        sort, column = query.get("sort", ""), query.get("column", "")
        if not isinstance(sort, str) or not isinstance(column, str):
            fail("INVALID_REQUEST", "ชื่อคอลัมน์สำหรับค้นหาหรือเรียงลำดับต้องเป็นข้อความ")
        if sort and sort not in columns or column and column not in columns:
            fail("INVALID_REQUEST", "ไม่พบคอลัมน์ที่ระบุ กรุณาเลือกคอลัมน์จากรายการ")
        keys = [column] if column else list(columns)
        where, params = "", []
        filters = query.get("filters") or []
        if filters:
            from dashboard import filter_clause
            semantics = query.get("columns") if isinstance(query.get("columns"), dict) else {}
            known = {key: {**value, **{field: semantics[key][field] for field in ("role", "time_format") if isinstance(semantics.get(key), dict) and field in semantics[key]}} for key, value in columns.items()}
            clause, filter_params = filter_clause(known, filters)
            where, params = f" WHERE ({clause})", filter_params
        if search:
            def search_expression(key):
                return f"CASE json_type(data, '$.{key}') WHEN 'true' THEN 'true' WHEN 'false' THEN 'false' ELSE casefold(json_extract(data, '$.{key}')) END"
            matches = " OR ".join(f"instr({search_expression(key)}, ?) > 0" for key in keys)
            where = f"{where} AND ({matches})" if where else f" WHERE ({matches})"
            params = [*params, *[search.casefold()] * len(keys)]
        order = "row_number ASC"
        if sort:
            expression = f"json_extract(data, '$.{sort}')"
            dtype = columns[sort]["data_type"]
            value = f"casefold({expression})" if dtype == "text" else f"julianday({expression})" if dtype == "date" else expression
            order = f"({expression} IS NULL) ASC, {value} {direction.upper()}, row_number ASC"
        table = f'"data_{sheet_id}"'
        total_rows = connection.execute(f"SELECT COUNT(*) FROM {table}{where}", params).fetchone()[0]
        items = connection.execute(f"SELECT row_number, data FROM {table}{where} ORDER BY {order} LIMIT ? OFFSET ?", [*params, page_size, (page - 1) * page_size]).fetchall()
        # Bound display payloads without modifying stored cells or the full
        # values used above for filtering/sorting. The UI must disclose this
        # preview-only abbreviation using the accompanying count and limit.
        max_cell_characters = min(2000, max(20, 1_000_000 // (page_size * len(columns) * 4)))
        truncated_cells = 0
        rows = []
        for row_number, stored in items:
            values = json.loads(stored)
            for key, value in values.items():
                if isinstance(value, str) and len(value) > max_cell_characters:
                    values[key] = value[:max_cell_characters - 1] + "…"
                    truncated_cells += 1
            rows.append({"row_number": row_number, "values": values})
        return {"rows": rows, "total_rows": total_rows, "page": page, "page_size": page_size,
                "truncated_cells": truncated_cells, "max_cell_characters": max_cell_characters}
    finally:
        connection.close()


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        if len(sys.argv) == 6 and sys.argv[1] == "ingest":
            result = ingest(sys.argv[2], sys.argv[3], sys.argv[4], json.loads(sys.argv[5]), lambda stage, progress: emit({"stage": stage, "progress": progress}))
        elif len(sys.argv) == 4 and sys.argv[1] == "preview":
            result = preview(sys.argv[2], json.loads(sys.argv[3]))
        elif len(sys.argv) == 4 and sys.argv[1] == "dashboard":
            from dashboard import run
            result = run(sys.argv[2], json.loads(Path(sys.argv[3]).read_text(encoding="utf-8")))
        elif len(sys.argv) == 3 and sys.argv[1] == "analyze":
            from analyzer import analyze
            result = analyze(sys.argv[2], lambda stage, progress: emit({"stage": stage, "progress": progress}))
        else:
            fail("INVALID_REQUEST", "รูปแบบคำสั่งไม่ถูกต้อง")
        emit({"result": result})
    except DatasetError as error:
        emit({"error": {"code": error.code, "message": error.message}})
        return 1
    except Exception:
        emit({"error": {"code": "PROCESSING_FAILED", "message": "ประมวลผลข้อมูลไม่สำเร็จ กรุณาลองใหม่"}})
        return 1
    return 0


if __name__ == "__main__":
    # Analyzer imports the shared protocol exception; keep one class identity
    # when this file is the CLI entrypoint rather than an imported module.
    sys.modules.setdefault("worker", sys.modules[__name__])
    sys.exit(main())
