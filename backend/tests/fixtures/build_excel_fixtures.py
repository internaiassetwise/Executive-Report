"""Build the spreadsheet test corpus in tests/fixtures/excel/.

Run from anywhere: python backend/tests/fixtures/build_excel_fixtures.py
The files are committed so tests never depend on this script, but rebuilding
must give the same content. Needs openpyxl and Pillow; xlwt only for legacy.xls.
Large files are not stored; tests generate them on the fly.

openpyxl saves formulas without a calculated result, while Excel always stores
one. `with_cached` writes the result Excel would have saved into the package so
the reader sees what it sees in a real file.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent / "excel"
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def save(book, name, cached=None, extra=None):
    buffer = io.BytesIO()
    book.save(buffer)
    data = buffer.getvalue()
    if cached or extra:
        data = rewrite(data, cached or {}, dict(extra or {}))
    (OUT / name).write_bytes(data)


def rewrite(data, cached, extra):
    """cached: {sheet part: {cell: value}}; extra: {part name: bytes | callable(old bytes) -> bytes}."""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename in cached:
                content = with_cached(content.decode("utf-8"), cached[entry.filename]).encode("utf-8")
            if entry.filename in extra:
                change = extra.pop(entry.filename)
                content = change(content) if callable(change) else change
            target.writestr(entry, content)
        for name, content in extra.items():
            target.writestr(name, content(b"") if callable(content) else content)
    return buffer.getvalue()


def with_cached(xml, values):
    for cell, value in values.items():
        error = isinstance(value, str) and value.startswith("#")
        pattern = re.compile(rf'<c r="{cell}"([^>]*)><f>(.*?)</f>(?:<v\s*/>|<v>[^<]*</v>)?</c>')
        match = pattern.search(xml)
        assert match, cell
        attrs = re.sub(r'\s*t="[^"]*"', "", match[1]) + (' t="e"' if error else ' t="str"' if isinstance(value, str) else "")
        xml = xml[:match.start()] + f'<c r="{cell}"{attrs}><f>{match[2]}</f><v>{value}</v></c>' + xml[match.end():]
    return xml


def simple():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["Order", "Region", "Product", "Amount", "Date"])
    rows = [("SO-001", "North", "Cement", 1200, date(2026, 1, 5)), ("SO-002", "South", "Sand", 800, date(2026, 1, 9)),
            ("SO-003", "North", "Steel", 4500, date(2026, 2, 2)), ("SO-004", "East", "Cement", 950, date(2026, 2, 14)),
            ("SO-005", "South", "Steel", 3100, date(2026, 3, 1)), ("SO-006", "East", "Sand", 600, date(2026, 3, 20))]
    for row in rows:
        sheet.append(row)
    save(book, "simple.xlsx")
    lines = ["Order,Region,Product,Amount,Date"] + [f"{o},{r},{p},{a},{d.isoformat()}" for o, r, p, a, d in rows]
    (OUT / "simple.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")


def multi_sheet():
    book = openpyxl.Workbook()
    book.remove(book.active)
    for month, factor in [("Jan", 1), ("Feb", 2), ("Mar", 3)]:
        sheet = book.create_sheet(month)
        sheet.append(["Item", "Qty", "Amount"])
        for item, qty in [("A", 1), ("B", 2), ("C", 3)]:
            sheet.append([item, qty * factor, qty * factor * 100])
    summary = book.create_sheet("Summary")
    summary.append(["Month", "Total"])
    for month, total in [("Jan", 600), ("Feb", 1200), ("Mar", 1800)]:
        summary.append([month, total])
    save(book, "multi_sheet.xlsx")


def formulas():
    book = openpyxl.Workbook()
    data = book.active
    data.title = "Data"
    data.append(["Item", "Qty", "Price", "Amount", "Share"])
    for row, (item, qty, price) in enumerate([("Pipe", 10, 25), ("Valve", 4, 150), ("Pump", 1, 0)], 2):
        data.append([item, qty, price, f"=B{row}*C{row}", f"=B{row}/C{row}"])
    data.append(["Total", "=SUM(B2:B4)", None, "=SUM(D2:D4)", None])
    summary = book.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Grand total", "=Data!D5"])
    summary.append(["Pipe amount", "='Data'!D2"])
    save(book, "formulas.xlsx", cached={
        "xl/worksheets/sheet1.xml": {"D2": 250, "E2": 0.4, "D3": 600, "E3": 0.0266666666666667, "D4": 0, "E4": "#DIV/0!", "B5": 15, "D5": 850},
        "xl/worksheets/sheet2.xml": {"B2": 850, "B3": 250}})


def merged_cells():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Merged"
    sheet.append(["Category", "Item", "Price", None])
    sheet.append([None, None, "Vendor A", "Vendor B"])
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:B2")
    sheet.merge_cells("C1:D1")
    for row in [["Structure", "Pile", 100, 110], [None, "Beam", 200, 190], [None, "Column", 300, 320],
                ["Finishing", "Paint", 50, 55], [None, "Tile", 80, 75]]:
        sheet.append(row)
    sheet.merge_cells("A3:A5")
    sheet.merge_cells("A6:A7")
    save(book, "merged_cells.xlsx")


def hidden_sheet():
    book = openpyxl.Workbook()
    visible = book.active
    visible.title = "Visible"
    visible.append(["Name", "Value", "Helper"])
    for index in range(1, 6):
        visible.append([f"Row {index}", index * 10, f"h{index}"])
    visible.row_dimensions[4].hidden = True
    visible.column_dimensions["C"].hidden = True
    hidden = book.create_sheet("Lookup")
    hidden.sheet_state = "hidden"
    hidden.append(["Code", "Label"])
    hidden.append(["A", "Alpha"])
    hidden.append(["B", "Beta"])
    save(book, "hidden_sheet.xlsx")


def multiple_tables():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Report"
    sheet.append(["Quantities"])
    sheet.append(["Item", "Qty", "Unit"])
    for row in [["Cement", 120, "bag"], ["Sand", 30, "m3"], ["Steel", 4, "ton"]]:
        sheet.append(row)
    sheet.append([])
    sheet.append(["Vendor prices"])
    sheet.append(["Vendor", "Cement", "Sand", "Steel"])
    for row in [["V1", 150, 400, 25000], ["V2", 145, 420, 24500]]:
        sheet.append(row)
    # A third table beyond the first 40 rows, found while reading.
    for _ in range(40):
        sheet.append([])
    sheet.append(["Payments"])
    sheet.append(["Date", "Paid"])
    sheet.append([date(2026, 4, 1), 50000])
    sheet.append([date(2026, 5, 1), 25000])
    save(book, "multiple_tables.xlsx")


def excel_tables():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet["A1"] = "Sales register (Excel table below)"
    rows = [["Region", "Rep", "Units", "Revenue"], ["North", "Ann", 10, 1000], ["South", "Bo", 7, 700], ["East", "Cy", 12, 1320],
            ["West", "Di", 5, 480], ["North", "Ed", 9, 990]]
    for offset, row in enumerate(rows):
        for column, value in enumerate(row):
            sheet.cell(row=3 + offset, column=2 + column, value=value)
    sheet.cell(row=9, column=2, value="Total")
    sheet.cell(row=9, column=5, value="=SUBTOTAL(109,Sales[Revenue])")
    table = Table(displayName="Sales", ref="B3:E9", totalsRowCount=1)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sheet.add_table(table)
    book.defined_names["Revenue"] = DefinedName("Revenue", attr_text="Data!$E$4:$E$8")
    sheet["C4"].comment = Comment("Ann covers two provinces", "QS")
    sheet["B4"].hyperlink = "https://example.com/regions/north"
    save(book, "excel_tables.xlsx", cached={"xl/worksheets/sheet1.xml": {"E9": 4490}})


def images():
    from PIL import Image, ImageDraw
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Site"
    sheet.append(["Zone", "Progress"])
    sheet.append(["A", 0.8])
    sheet.append(["B", 0.45])
    picture = Image.new("RGB", (360, 160), "white")
    draw = ImageDraw.Draw(picture)
    for index, line in enumerate(["Item     Qty   Price", "Cement   120   150", "Sand      30   400"]):
        draw.text((14, 16 + index * 40), line, fill="black")
    draw.rectangle([4, 4, 355, 155], outline="black")
    buffer = io.BytesIO()
    picture.save(buffer, "PNG")
    sheet.add_image(openpyxl.drawing.image.Image(io.BytesIO(buffer.getvalue())), "D2")
    logo = Image.new("RGB", (120, 60), (20, 80, 160))
    buffer = io.BytesIO()
    logo.save(buffer, "JPEG")
    sheet.add_image(openpyxl.drawing.image.Image(io.BytesIO(buffer.getvalue())), "D12")
    save(book, "images.xlsx")


def charts():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Monthly"
    sheet.append(["Month", "Revenue", "Cost"])
    for row in [["Jan", 100, 80], ["Feb", 120, 90], ["Mar", 90, 70], ["Apr", 150, 100]]:
        sheet.append(row)
    bar = BarChart()
    bar.title = "Revenue by month"
    bar.add_data(Reference(sheet, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    bar.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=5))
    sheet.add_chart(bar, "E2")
    line = LineChart()
    line.title = "Cost trend"
    line.add_data(Reference(sheet, min_col=3, min_row=1, max_row=5), titles_from_data=True)
    line.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=5))
    sheet.add_chart(line, "E20")
    save(book, "charts.xlsx")


def pivot():
    book = openpyxl.Workbook()
    data = book.active
    data.title = "Data"
    data.append(["Region", "Product", "Amount"])
    for row in [["North", "A", 100], ["North", "B", 50], ["South", "A", 70], ["South", "B", 30], ["East", "A", 40]]:
        data.append(row)
    report = book.create_sheet("Pivot")
    report["A1"] = "Sum of Amount by Region"
    report.append([])
    report.append(["Row Labels", "Sum of Amount"])
    for row in [["East", 40], ["North", 150], ["South", 100], ["Grand Total", 290]]:
        report.append(row)
    pivot_xml = (f'<pivotTableDefinition xmlns="{MAIN}" name="RegionPivot" cacheId="1" dataCaption="Values">'
                 '<location ref="A3:B7" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/></pivotTableDefinition>').encode()
    cache_xml = (f'<pivotCacheDefinition xmlns="{MAIN}" xmlns:r="{REL}" recordCount="0"><cacheSource type="worksheet">'
                 '<worksheetSource ref="A1:C6" sheet="Data"/></cacheSource><cacheFields count="0"/></pivotCacheDefinition>').encode()

    def rels(target, kind):
        return (f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{PKG}"><Relationship Id="rId1" '
                f'Type="{REL}/{kind}" Target="{target}"/></Relationships>').encode()

    def types(old):
        overrides = ('<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
                     '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>')
        return old.replace(b"</Types>", overrides.encode() + b"</Types>")

    save(book, "pivot.xlsx", extra={
        "[Content_Types].xml": types,
        "xl/pivotTables/pivotTable1.xml": pivot_xml,
        "xl/pivotTables/_rels/pivotTable1.xml.rels": rels("../pivotCache/pivotCacheDefinition1.xml", "pivotCacheDefinition"),
        "xl/pivotCache/pivotCacheDefinition1.xml": cache_xml,
        "xl/worksheets/_rels/sheet2.xml.rels": rels("../pivotTables/pivotTable1.xml", "pivotTable"),
    })


def thai():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "ค่าใช้จ่าย"
    sheet.append(["วันที่", "หมวด", "รายละเอียด", "จำนวนเงิน", "สัดส่วน"])
    rows = [(date(2026, 1, 3), "ค่าน้ำ", "ค่าน้ำประปา สำนักงาน", 1250.5, 0.1), (date(2026, 1, 5), "ค่าไฟ", "ค่าไฟฟ้า อาคาร ก", 8420, 0.6),
            (date(2026, 1, 9), "ค่าเดินทาง", "ค่าทางด่วน 🚗", 560, 0.05), (date(2026, 2, 1), "ค่าน้ำ", "ค่าน้ำประปา สำนักงาน", 1320, 0.1),
            (date(2026, 2, 3), "ค่าไฟ", "ค่าไฟฟ้า อาคาร ก", 7950, 0.15)]
    for row, values in enumerate(rows, 2):
        sheet.append(values)
        sheet.cell(row=row, column=4).number_format = '#,##0.00 "฿"'
        sheet.cell(row=row, column=5).number_format = "0.0%"
    save(book, "thai.xlsx")
    lines = ["หมวด,จำนวนเงิน", "ค่าน้ำ,1250.50", "ค่าไฟ,\"8,420\""]
    (OUT / "thai_cp874.csv").write_bytes(("\r\n".join(lines) + "\r\n").encode("cp874"))


def messy_real_world():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "BOQ ตึก A"
    rows = [
        ["บริษัท ตัวอย่าง คอนสตรัคชั่น จำกัด"],
        ["ใบเสนอราคา / BOQ โครงการอาคาร A"],
        ["วันที่ 15/08/2569"],
        [],
        ["ลำดับ", "หมวดงาน", "รายการ", "จำนวน", "ราคา (บาท)", None, "helper"],
        [None, None, None, None, "ต่อหน่วย", "รวม"],
        [1, "งานโครงสร้าง", "เสาเข็ม", 20, "1,250.00", 25000, "x"],
        [2, None, "ฐานราก", 10, 3000, "=D8*E8"],
        [3, None, "คาน", 15, 2000, "=D9*E9"],
        [None, "รวมหมวดงานโครงสร้าง", None, None, None, "=SUM(F7:F9)"],
        [],
        ["ลำดับ", "หมวดงาน", "รายการ", "จำนวน", "ราคา (บาท)"],
        [None, None, None, None, "ต่อหน่วย", "รวม"],
        [4, "งานสถาปัตย์", "ผนังก่ออิฐ", 100, 350, 35000],
        [5, None, "ฉาบปูน", 200, 150, 30000],
        ["รวมทั้งสิ้น", None, None, None, None, 150000],
        ["หมายเหตุ: ราคารวมภาษีมูลค่าเพิ่ม 7% แล้ว"],
        ["ลงชื่อ", "............", "ผู้จัดทำ"],
        [],
        ["สรุปตามหมวดงาน"],
        ["หมวดงาน", "มูลค่า (บาท)", "สัดส่วน"],
        ["งานโครงสร้าง", 85000, 0.567],
        ["งานสถาปัตย์", 65000, 0.433],
    ]
    for row in rows:
        sheet.append(row)
    for header in (5, 12):
        for column in "ABCD":
            sheet.merge_cells(f"{column}{header}:{column}{header + 1}")
        sheet.merge_cells(f"E{header}:F{header}")
    sheet.merge_cells("A1:F1")
    sheet.merge_cells("A2:F2")
    sheet.merge_cells("B7:B9")
    sheet.merge_cells("B14:B15")
    sheet.column_dimensions["G"].hidden = True
    for row in (22, 23):
        sheet.cell(row=row, column=3).number_format = "0.0%"
    summary = book.create_sheet("สรุป")
    summary.append(["รายการ", "มูลค่า"])
    summary.append(["มูลค่ารวม", "='BOQ ตึก A'!F16"])
    save(book, "messy_real_world.xlsx", cached={"xl/worksheets/sheet1.xml": {"F8": 30000, "F9": 30000, "F10": 85000},
                                                "xl/worksheets/sheet2.xml": {"B2": 150000}})


def legacy_xls():
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet("Legacy")
    header = xlwt.easyxf("font: bold on")
    money = xlwt.easyxf(num_format_str="#,##0.00")
    for column, name in enumerate(["Code", "Group", "Amount"]):
        sheet.write(0, column, name, header)
    for row, (code, amount) in enumerate([("P1", 1200.5), ("P2", 800), ("P3", 450)], 1):
        sheet.write(row, 0, code)
        sheet.write(row, 2, amount, money)
    sheet.write_merge(1, 2, 1, 1, "Civil")
    sheet.write(3, 1, "MEP")
    sheet.row(3).hidden = True
    book.save(str(OUT / "legacy.xls"))


def corrupted():
    data = (OUT / "simple.xlsx").read_bytes()
    (OUT / "corrupted.xlsx").write_bytes(data[: len(data) * 3 // 5])


def main():
    OUT.mkdir(exist_ok=True)
    for build in (simple, multi_sheet, formulas, merged_cells, hidden_sheet, multiple_tables, excel_tables, images, charts, pivot, thai,
                  messy_real_world, legacy_xls, corrupted):
        build()
    print("\n".join(sorted(path.name for path in OUT.iterdir())))


if __name__ == "__main__":
    main()
