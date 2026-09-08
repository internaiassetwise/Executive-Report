import io, json, sys, unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'analysis'))
import analysis_engine as e
import boq_engine as B
import boq_report as R
import openpyxl


def workbook_bytes(book):
    stream = io.BytesIO(); book.save(stream); return stream.getvalue()


def sheets_of(book):
    return [{'name': ws.title, 'grid': [list(r) for r in ws.iter_rows(values_only=True)]} for ws in book.worksheets]


# Twelve benchmark lines. Vendor AAA prices material 2x the benchmark on the
# first line only; vendor BBB pads every quantity by 20% and labour by 5%.
ITEMS = [(f'Item {i}', 'หมวด ' + ('A' if i < 6 else 'B'), 10.0 * (i + 1), 100.0, 50.0) for i in range(12)]


def side_by_side_book():
    b = openpyxl.Workbook(); b.remove(b.active)
    for sheet in ('ST_A', 'AR_A'):
        s = b.create_sheet(sheet)
        s.append(['No', 'รายการ', 'หมวดงาน', 'หน่วย', 'ปริมาณ RBP', 'ปริมาณ AAA เสนอ', 'ปริมาณ BBB เสนอ',
                  'ราคาของ RBP', 'ราคาของ AAA เสนอ', 'ราคาของ BBB เสนอ', 'ราคาแรง RBP', 'ราคาแรง AAA เสนอ', 'ราคาแรง BBB เสนอ',
                  'ส่วนต่าง AAA', 'Flag ปริมาณ'])
        for i, (item, cat, q, m, l) in enumerate(ITEMS, 1):
            s.append([i, item, cat, 'ea', q, q, q * 1.2, m, m * (2 if i == 1 else 1), m, l, l, l * 1.05, 0, None])
    return b


def sheet_per_vendor_book():
    b = openpyxl.Workbook(); b.remove(b.active)
    for vendor, qf, mf, lf in (('AAA', 1.0, 1.5, 1.0), ('BBB', 1.3, 1.0, 1.0)):
        s = b.create_sheet(vendor)
        s.append(['No', 'รายการ', 'หมวดงาน', 'หน่วย', 'ปริมาณ (ราคากลาง)', 'ปริมาณ', 'ราคาวัสดุ/หน่วย (ราคากลาง)',
                  'ราคาวัสดุ/หน่วย', 'ราคาแรง/หน่วย (ราคากลาง)', 'ราคาแรง/หน่วย'])
        for i, (item, cat, q, m, l) in enumerate(ITEMS, 1):
            s.append([i, item, cat, 'ea', q, q * qf, m, m * mf, l, l * lf])
    return b


def single_vendor_book(label_cell=None):
    b = openpyxl.Workbook(); s = b.active; s.title = 'ST_A'
    if label_cell:
        s.append([label_cell]); s.append([])
    s.append(['No', 'รายการ', 'ปริมาณ (ราคากลาง)', 'ปริมาณ', 'ราคาวัสดุ/หน่วย (ราคากลาง)', 'ราคาวัสดุ/หน่วย',
              'ราคาแรง/หน่วย (ราคากลาง)', 'ราคาแรง/หน่วย'])
    for i, (item, _, q, m, l) in enumerate(ITEMS, 1):
        s.append([i, item, q, q, m, m * 1.5, l, l])
    return b


class Detection(unittest.TestCase):
    def test_side_by_side_vendors_in_one_sheet(self):
        rep = B.build_many([(sheets_of(side_by_side_book()), 'compare.xlsx')])
        self.assertEqual([v['vendor'] for v in rep['vendors']], ['AAA', 'BBB'])
        self.assertEqual(rep['groups'], ['AR', 'ST'])
        self.assertEqual(rep['tolerance_source'], 'default')
        aaa, bbb = rep['vendors']
        self.assertEqual(aaa['benchmark'], 'RBP')
        # AAA: one material line at 2x on each sheet; nothing else over.
        self.assertEqual(aaa['total']['material_over'], 2)
        self.assertEqual(aaa['total']['labour_over'], 0)
        self.assertEqual(aaa['total']['quantity_over'], 0)
        # BBB: every quantity 20% over the default 15% tolerance; labour +5% is within it.
        self.assertEqual(bbb['total']['quantity_over'], 24)
        self.assertEqual(bbb['total']['labour_over'], 0)
        self.assertAlmostEqual(bbb['total']['labour_dev_pct'], 5.0, places=6)
        self.assertAlmostEqual(bbb['total']['material_dev_pct'], 0.0, places=6)
        # Derived columns never become inputs.
        self.assertNotIn('ส่วนต่าง', json.dumps(rep, ensure_ascii=False))
        # Normalization is one-directional and priced from the unit rates.
        self.assertAlmostEqual(aaa['total']['savings'], 2 * 10.0 * 100.0, places=6)
        self.assertGreater(bbb['total']['savings'], 0)
        self.assertEqual(len(rep['comparison']['quantity_over']['ST']), 2)

    def test_one_sheet_per_vendor(self):
        rep = B.build_many([(sheets_of(sheet_per_vendor_book()), 'tabs.xlsx')])
        self.assertEqual([v['vendor'] for v in rep['vendors']], ['AAA', 'BBB'])
        self.assertEqual(rep['groups'], ['หมวด A', 'หมวด B'])
        aaa, bbb = rep['vendors']
        self.assertEqual(aaa['total']['material_over'], 12)
        self.assertEqual(bbb['total']['quantity_over'], 12)
        self.assertEqual(aaa['sheets_used'], ['AAA'])

    def test_single_unlabelled_vendor_is_named_from_the_file(self):
        rep = B.build_many([(sheets_of(single_vendor_book()), 'BOQ_Wisawapat.xlsx')])
        self.assertEqual(rep['vendors'][0]['vendor'], 'BOQ_Wisawapat')
        rep = B.build_many([(sheets_of(single_vendor_book('ผู้รับเหมา: วิศวพัฒน์ จำกัด')), 'x.xlsx')])
        self.assertEqual(rep['vendors'][0]['vendor'], 'วิศวพัฒน์ จำกัด')
        self.assertEqual(rep['vendors'][0]['benchmark'], 'ราคากลาง')

    def test_same_vendor_in_two_files_stays_distinguishable(self):
        rep = B.build_many([(sheets_of(single_vendor_book()), 'a.xlsx'), (sheets_of(single_vendor_book()), 'b.xlsx')])
        self.assertEqual(len(rep['vendors']), 2)
        self.assertNotEqual(rep['vendors'][0]['vendor'], rep['vendors'][1]['vendor'])

    def test_workbook_without_benchmark_columns_is_not_boq(self):
        b = openpyxl.Workbook(); s = b.active
        s.append(['Region', 'Amount']); s.append(['North', 10]); s.append(['South', 20])
        self.assertIsNone(B.build_many([(sheets_of(b), 'sales.xlsx')]))


class Dispatch(unittest.TestCase):
    def test_boq_action_returns_report_and_html(self):
        progress = []
        out = e.dispatch('boq', {'files': [{'filename': 'compare.xlsx', 'bytes': workbook_bytes(side_by_side_book())}]},
                         lambda d, t, n: progress.append((d, t, n)))
        self.assertEqual(out['mode'], 'boq')
        self.assertEqual(progress, [(1, 1, 'compare.xlsx')])
        html = out['html']
        for heading in ('บทนำและขอบเขตการวิเคราะห์', 'ผู้เสนองาน: AAA', 'ผู้เสนองาน: BBB', 'ตารางที่ 1', 'ตารางที่ 2', 'ตารางที่ 3',
                        'การวิเคราะห์เชิงลึก', 'การวิเคราะห์เปรียบเทียบภาพรวมทุกเจ้า', 'Signature Pattern', 'บทวิเคราะห์เชิงกลยุทธ์',
                        'บทสรุปผู้บริหาร', 'Grand Total Project Cost'):
            self.assertIn(heading, html)
        json.dumps(out, ensure_ascii=False, allow_nan=False)
        rebuilt = e.dispatch('boq_rebuild', {'tolerance': 0.01})
        self.assertEqual(rebuilt['report']['tolerance_source'], 'declared')
        self.assertEqual(rebuilt['report']['vendors'][1]['total']['labour_over'], 24)

    def test_single_generic_upload_falls_through_with_its_profile(self):
        b = openpyxl.Workbook(); s = b.active
        s.append(['Region', 'Amount']); s.append(['North', 10]); s.append(['South', 20])
        out = e.dispatch('boq', {'files': [{'filename': 'sales.xlsx', 'bytes': workbook_bytes(b)}]})
        self.assertEqual(out['mode'], 'generic')
        self.assertEqual(out['book']['tables_count'], 1)
        self.assertEqual(e.dispatch('boq', {'files': [{'filename': 'a.xlsx', 'bytes': workbook_bytes(b)},
                                                      {'filename': 'b.xlsx', 'bytes': workbook_bytes(b)}]})['mode'], 'none')

    def test_single_vendor_report_omits_the_comparison_section(self):
        html = R.render(B.build_many([(sheets_of(single_vendor_book()), 'one.xlsx')]))
        self.assertIn('ผู้เสนองาน: one', html)
        self.assertNotIn('การวิเคราะห์เปรียบเทียบภาพรวมทุกเจ้า', html)
        self.assertIn('บทสรุปผู้บริหาร', html)


if __name__ == '__main__':
    unittest.main()
