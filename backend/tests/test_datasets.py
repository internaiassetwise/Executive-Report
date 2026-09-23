import json
import io
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "datasets"))
from worker import DatasetError, ingest, preview


class DatasetTests(unittest.TestCase):
    def setUp(self):
        self.previous_tempdir = tempfile.tempdir
        self.workspace = Path(__file__).resolve().parents[2]
        # Windows Python 3.12's mode-0700 temporary directories exclude the
        # restricted token used by local sandboxes. Inherit workspace access.
        self.root = self.workspace / f"dataset-test-{uuid.uuid4().hex}"
        self.root.mkdir(mode=0o777)
        tempfile.tempdir = str(self.root)
        self.database = self.root / "dataset.sqlite"

    def tearDown(self):
        tempfile.tempdir = self.previous_tempdir
        self.assertEqual(self.root.resolve().parent, self.workspace)
        shutil.rmtree(self.root)

    def csv(self, text, encoding="utf-8", limits=None):
        source = self.root / "input.csv"
        source.write_bytes(text.encode(encoding))
        return ingest(source, self.database, "input.csv", limits)

    def workbook(self, book, limits=None):
        source = self.root / "input.xlsx"
        book.save(source)
        return ingest(source, self.database, "input.xlsx", limits)

    def error(self, code, callback):
        with self.assertRaises(DatasetError) as caught:
            callback()
        self.assertEqual(caught.exception.code, code)
        self.assertFalse(self.database.exists(), "Failed ingestion must clean up its database")

    def test_csv_types_original_row_numbers_and_duplicates(self):
        result = self.csv("Code,Value,When,Flag,Empty\n001,2,2026-01-02,true,\n\n002,10,2026-01-03,false,\n002,10,2026-01-03,false,\n")
        self.assertEqual(result["rows_count"], 3)
        self.assertEqual([c["data_type"] for c in result["sheets"][0]["columns"]], ["text", "number", "date", "boolean", "empty"])
        rows = preview(self.database, {})["rows"]
        self.assertEqual([row["row_number"] for row in rows], [2, 4, 5])
        self.assertEqual(rows[0]["values"], {"c0": "001", "c1": 2, "c2": "2026-01-02", "c3": True, "c4": None})
        self.assertEqual(preview(self.database, {"column": "c3", "search": "true"})["total_rows"], 1)

    def test_utf16_semicolon_and_quoted_newlines(self):
        result = self.csv('ชื่อ;จำนวน\n"หนึ่ง\nสอง";5\nสาม;8', "utf-16")
        self.assertEqual(result["rows_count"], 2)
        self.assertEqual(preview(self.database, {})["rows"][0]["values"]["c0"], "หนึ่ง\nสอง")
        self.assertEqual(preview(self.database, {})["rows"][0]["row_number"], 2)

    def test_date_sort_uses_instants_and_keeps_nulls_last(self):
        self.csv("Name,When\nLater,2026-01-01T01:00:00+00:00\nEarlier,2026-01-01T02:00:00+07:00\nUnknown,")
        rows = preview(self.database, {"sort": "c1"})["rows"]
        self.assertEqual([row["values"]["c0"] for row in rows], ["Earlier", "Later", "Unknown"])

    def test_preview_abbreviation_preserves_storage_tail_search_and_sort(self):
        first = "x" * 4000 + "a-tail-match"
        second = "x" * 4000 + "z-tail-match"
        self.csv(f"Name,Comment\nLater,{second}\nEarlier,{first}\n")
        result = preview(self.database, {"search": "tail-match", "column": "c1", "sort": "c1"})
        self.assertEqual(result["total_rows"], 2)
        self.assertEqual(result["truncated_cells"], 2)
        self.assertEqual(result["max_cell_characters"], 2000)
        self.assertEqual([row["values"]["c0"] for row in result["rows"]], ["Earlier", "Later"])
        self.assertEqual(len(result["rows"][0]["values"]["c1"]), 2000)
        self.assertTrue(result["rows"][0]["values"]["c1"].endswith("…"))
        exact_tail = preview(self.database, {"search": "z-tail-match", "column": "c1"})
        self.assertEqual(exact_tail["total_rows"], 1)
        connection = sqlite3.connect(self.database)
        try:
            stored = json.loads(connection.execute('SELECT data FROM "data_s0" WHERE row_number=2').fetchone()[0])
        finally:
            connection.close()
        self.assertEqual(stored["c1"], second)

    def test_wide_unicode_previews_bound_payload_for_every_page_size(self):
        headers = [f"Column{index}" for index in range(200)]
        value = "😀" * 250
        text = ",".join(headers) + "\n" + "\n".join(",".join([value] * 200) for _ in range(100))
        self.csv(text)
        for page_size in [1, 25, 50, 100]:
            with self.subTest(page_size=page_size):
                result = preview(self.database, {"page_size": page_size})
                limit = min(2000, max(20, 1_000_000 // (page_size * 200 * 4)))
                self.assertEqual(result["max_cell_characters"], limit)
                self.assertEqual(result["truncated_cells"], page_size * 200 if len(value) > limit else 0)
                self.assertTrue(all(len(cell) <= limit for row in result["rows"] for cell in row["values"].values()))
                self.assertLess(len(json.dumps(result, ensure_ascii=False).encode("utf-8")), 2_000_000)

    def test_sort_search_column_filter_and_pagination(self):
        self.csv("Name,Value\nAlpha,2\nbeta,10\nALPHA,1\nGamma,\nLiteral%_,3")
        page = preview(self.database, {"sort": "c1", "direction": "desc", "page_size": 2, "page": 2})
        self.assertEqual([r["values"]["c1"] for r in page["rows"]], [2, 1])
        self.assertEqual(page["total_rows"], 5)
        last = preview(self.database, {"sort": "c1", "direction": "desc", "page_size": 2, "page": 3})
        self.assertIsNone(last["rows"][0]["values"]["c1"])
        self.assertEqual(preview(self.database, {"search": "alpha", "column": "c0"})["total_rows"], 2)
        self.assertEqual(preview(self.database, {"search": "alpha", "column": "c1"})["total_rows"], 0)
        self.assertEqual(preview(self.database, {"search": "%_"})["total_rows"], 1)
        for query in [{"sort": "c0); DROP TABLE metadata;--"}, {"sheet": "s99"}, {"page_size": 101}, {"page": 0}, {"direction": "random"}]:
            with self.subTest(query=query), self.assertRaises(DatasetError):
                preview(self.database, query)

    def test_invalid_headers_and_corrupted_csv_fail(self):
        cases = [("Value,Value\n1,2", "DUPLICATE_HEADERS"), (",Value\na,2", "MISSING_HEADERS"), ("1,2\n3,4", "MISSING_HEADERS"), ('Name,Value\nA,"1\nB,2', "INVALID_FILE"), ("Name,Value\nA,2,3", "MISSING_HEADERS"), ("Name,Value\n", "EMPTY_DATASET"), ("", "EMPTY_FILE")]
        for value, code in cases:
            with self.subTest(value=value):
                self.error(code, lambda: self.csv(value))

    def test_bounded_total_rows_cells_and_columns(self):
        for limits in [{"max_rows": 1}, {"max_cells": 3}, {"max_columns": 1}]:
            with self.subTest(limits=limits):
                self.error("LIMIT_EXCEEDED", lambda: self.csv("A,B\n1,2\n3,4", limits=limits))
        self.assertEqual(self.csv("A,B\n1,2\n3,4", limits={"max_rows": 2, "max_cells": 4, "max_columns": 2})["rows_count"], 2)

    def test_binary_cells_and_invalid_encoding_fail_without_data_in_error(self):
        self.error("UNSUPPORTED_VALUE", lambda: self.csv("Name,Value\nsecret\x00content,2"))
        source = self.root / "bad.csv"
        source.write_bytes(b"Name,Value\n\xff,2")
        self.error("INVALID_ENCODING", lambda: ingest(source, self.database, "bad.csv"))

    def test_xlsx_all_sheets_formulas_dates_hidden_and_totals_preserved(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "BOQ"
        sheet.append(["Item", "Amount", "When"])
        sheet.append(["A", 10, datetime(2026, 1, 2)])
        sheet.append(["B", 20, datetime(2026, 1, 3)])
        sheet.append(["Total", "=SUM(B2:B3)", None])
        hidden = book.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden.append(["Code", "Flag"])
        hidden.append(["0001", True])
        book.create_sheet("Empty")
        result = self.workbook(book)
        self.assertEqual(result["rows_count"], 4)
        self.assertEqual(result["columns_count"], 5)
        self.assertEqual([s["name"] for s in result["sheets"]], ["BOQ", "Hidden"])
        self.assertTrue(result["warnings"])
        self.assertTrue(result["sheets"][0]["warnings"])
        rows = preview(self.database, {"sheet": "s0"})["rows"]
        self.assertEqual(rows[-1]["values"]["c0"], "Total")
        # openpyxl saves no calculated result, so the formula cell is empty and
        # disclosed; formula text never becomes a data value.
        self.assertIsNone(rows[-1]["values"]["c1"])
        warnings = " ".join(result["sheets"][0]["warnings"])
        self.assertIn("ไม่มีค่าที่คำนวณไว้", warnings)
        self.assertIn("แถวสรุปยอด", warnings)
        self.assertTrue(preview(self.database, {"sheet": "s1"})["rows"][0]["values"]["c1"])

    def test_xlsx_invalid_sheet_rolls_back_other_sheets(self):
        book = openpyxl.Workbook()
        book.active.append(["Good"])
        book.active.append([1])
        bad = book.create_sheet("Bad")
        bad.append(["Duplicate", "Duplicate"])
        bad.append([1, 2])
        self.error("DUPLICATE_HEADERS", lambda: self.workbook(book))

    def test_xlsx_limits_and_corruption(self):
        book = openpyxl.Workbook()
        book.active.append(["A"])
        book.active.append([1])
        second = book.create_sheet("Second")
        second.append(["B"])
        second.append([2])
        for limits in [{"max_sheets": 1}, {"max_rows": 1}, {"max_uncompressed_bytes": 5}]:
            with self.subTest(limits=limits):
                self.error("LIMIT_EXCEEDED", lambda: self.workbook(book, limits))
        source = self.root / "invalid.xlsx"
        source.write_bytes(b"not a zip")
        self.error("INVALID_FILE", lambda: ingest(source, self.database, "invalid.xlsx"))

    def test_macro_payload_is_rejected(self):
        source = self.root / "macro.xlsx"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("xl/workbook.xml", "<workbook/>")
            archive.writestr("xl/vbaProject.bin", b"not executed")
        self.error("UNSUPPORTED_FORMAT", lambda: ingest(source, self.database, "macro.xlsx"))

    def test_corrupt_xlsx_coordinates_do_not_silently_drop_rows(self):
        book = openpyxl.Workbook()
        book.active.append(["Value"])
        book.active.append([11])
        book.active.append([22])
        buffer = io.BytesIO()
        book.save(buffer)
        for replacement in [b'<row r="2"><c r="A2"', b'<row r="1000000000"><c r="A1000000000"']:
            with self.subTest(replacement=replacement):
                source = self.root / "corrupt.xlsx"
                with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as original, zipfile.ZipFile(source, "w") as modified:
                    for entry in original.infolist():
                        data = original.read(entry.filename)
                        if entry.filename == "xl/worksheets/sheet1.xml":
                            data = data.replace(b'<row r="3"><c r="A3"', replacement)
                        modified.writestr(entry, data)
                self.error("INVALID_FILE", lambda: ingest(source, self.database, "corrupt.xlsx"))

    def test_cli_protocol_contains_progress_and_metadata_only(self):
        source = self.root / "input.csv"
        source.write_text("Name,Value\nSENSITIVE_UNIQUE_VALUE,2", encoding="utf-8")
        worker = Path(__file__).resolve().parents[1] / "datasets" / "worker.py"
        process = subprocess.run([sys.executable, "-B", "-u", str(worker), "ingest", str(source), str(self.database), "input.csv", "{}"], capture_output=True, encoding="utf-8")
        self.assertEqual(process.returncode, 0, process.stdout)
        self.assertNotIn("SENSITIVE_UNIQUE_VALUE", process.stdout + process.stderr)
        events = [json.loads(line) for line in process.stdout.splitlines()]
        self.assertEqual(events[-1]["result"]["rows_count"], 1)
        self.assertEqual([event["progress"] for event in events[:-1]], sorted(event["progress"] for event in events[:-1]))

    def test_missing_xlsx_dependency_has_actionable_error(self):
        book = openpyxl.Workbook()
        book.active.append(["Value"])
        book.active.append([1])
        source = self.root / "input.xlsx"
        book.save(source)
        worker = Path(__file__).resolve().parents[1] / "datasets" / "worker.py"
        # -S excludes site-packages, reproducing a fresh Python installation.
        process = subprocess.run([sys.executable, "-S", "-B", str(worker), "ingest", str(source), str(self.database), "input.xlsx", "{}"], capture_output=True, encoding="utf-8")
        self.assertEqual(process.returncode, 1)
        self.assertEqual(json.loads(process.stdout.splitlines()[-1])["error"]["code"], "PARSER_UNAVAILABLE")
        self.assertEqual(process.stderr, "")
        self.assertFalse(self.database.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
