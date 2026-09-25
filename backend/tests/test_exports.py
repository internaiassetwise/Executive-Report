import csv
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import uuid
from pathlib import Path

import openpyxl
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "datasets"))
from exports import ExportError, export_report
from worker import ingest


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.workspace = Path(__file__).resolve().parents[2]
        self.root = self.workspace / f"export-test-{uuid.uuid4().hex}"
        self.root.mkdir(mode=0o777)
        self.previous_tempdir = tempfile.tempdir
        tempfile.tempdir = str(self.root)
        self.database = self.root / "dataset.sqlite"
        source = self.root / "source.csv"
        with source.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["ชื่อ", "จำนวน"])
            writer.writerows([["ปูนซีเมนต์", 10], ["  =HYPERLINK(\"bad\")", -20], ["\t@SUM(1)", 3], ["-not a number", 4]])
        ingest(source, self.database, "ข้อมูลต้นทาง.csv")
        self.analysis_path = self.root / "analysis.json"
        self.analysis = {
            "generated_at": "2026-09-10T08:00:00Z", "summary": "รายงานสรุปจากข้อมูลต้นทางสำหรับผู้บริหาร",
            "kpis": [{"id": "k1", "name": "จำนวนแถว", "value": 4, "formatted_value": "4", "source": {"sheet": "CSV"}, "method": "นับทุกแถวข้อมูล"}],
            "insights": [{"id": "EV-001", "title": "ข้อมูลจำนวนแถว", "description": "มีข้อมูลทั้งหมด 4 แถว", "evidence": {"metric": "rows", "value": 4, "sheet": "CSV", "method": "นับแถว", "columns": []}}],
            "profiles": [{"sheet_id": "s0", "sheet_name": "CSV", "rows_count": 4, "duplicate_rows": 0, "missing_count": 0, "missing_percentage": 0, "columns": [], "warnings": []}],
            "charts": [{"id": "chart1", "type": "bar", "title": "จำนวนตามรายการ", "x_label": "รายการ", "y_label": "จำนวน", "data": [{"x": "ปูนซีเมนต์", "y": 10}, {"x": "อื่น ๆ", "y": -13}], "method": "ผลรวมจำนวน"}],
            "report": {"sections": [{"id": f"section-{index}", "title": f"{index}. ข้อค้นพบจากข้อมูล", "paragraphs": ["ผลการวิเคราะห์สถิติและคุณภาพข้อมูล เพื่อประกอบการทบทวนของผู้บริหาร โดยเก็บสูตรเป็นข้อความและไม่คำนวณซ้ำ"], "evidence_ids": ["EV-001"]} for index in range(1, 11)]},
        }
        self.analysis_path.write_text(json.dumps(self.analysis, ensure_ascii=False), encoding="utf-8")

    def tearDown(self):
        tempfile.tempdir = self.previous_tempdir
        self.assertEqual(self.root.resolve().parent, self.workspace)
        shutil.rmtree(self.root)

    def test_csv_all_records_utf8_and_formula_safety_without_analysis(self):
        self.analysis_path.write_text("{}", encoding="utf-8")
        output = self.root / "report.csv"
        export_report(self.database, self.analysis_path, "csv", output, "s0")
        self.assertTrue(output.read_bytes().startswith(b"\xef\xbb\xbf"))
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ["ชื่อ", "จำนวน"])
        self.assertEqual(rows[1], ["ปูนซีเมนต์", "10"])
        self.assertEqual(rows[2][1], "-20")
        self.assertTrue(all(rows[index][0].startswith("'") for index in (2, 3, 4)))
        with self.assertRaises(ExportError):
            export_report(self.database, self.analysis_path, "csv", output, 's0";DROP TABLE metadata;--')

    def test_xlsx_native_data_and_formula_text(self):
        output = self.root / "report.xlsx"
        export_report(self.database, self.analysis_path, "xlsx", output)
        book = openpyxl.load_workbook(output, data_only=False)
        try:
            self.assertEqual(book.sheetnames[:6], ["Overview", "KPIs", "Insights", "Quality", "Columns", "Charts"])
            self.assertEqual(book["KPIs"]["B2"].value, 4)
            self.assertEqual(len(book["Charts"]._charts), 1)
            data = book["Data 1 CSV"]
            self.assertEqual(data["B3"].value, -20)
            self.assertEqual(data["A2"].value, "ปูนซีเมนต์")
            self.assertEqual(data["A3"].value, '  =HYPERLINK("bad")')
            for index in (3, 4, 5):
                self.assertEqual(data.cell(index, 1).data_type, "s")
                self.assertTrue(data.cell(index, 1).quotePrefix)
        finally:
            book.close()

    def test_pdf_is_typeset_readable_and_cli_returns_no_data(self):
        output = self.root / "report.pdf"
        script = Path(__file__).resolve().parents[1] / "datasets" / "exports.py"
        result = subprocess.run([sys.executable, "-B", str(script), "export", str(self.database), str(self.analysis_path), "pdf", str(output)], capture_output=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(json.loads(result.stdout)["result"]["path"], str(output.resolve()))
        self.assertNotIn("HYPERLINK", result.stdout + result.stderr)
        self.assertTrue(output.read_bytes().startswith(b"%PDF-"))
        reader = PdfReader(output)
        self.assertGreaterEqual(len(reader.pages), 2)
        content = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("ASW Data Insight", content)
        self.assertIn("CSV | rows: 4", content, "the evidence appendix is included")
        self.assertNotIn("EV-001", content, "evidence ids are for checking, not for readers")
        self.assertIn("รายงาน", content)
        self.assertFalse(output.with_name(output.name + ".part").exists())


    def test_pdf_has_the_pages_asked_for_and_cited_tables(self):
        insight = self.analysis["insights"][0]
        self.analysis["insights"] = [dict(insight, id=f"EV-{index:03d}") for index in range(1, 80)]
        self.analysis["report"]["source"] = "ai"
        self.analysis["report"]["tables"] = {"Q-001": {"id": "Q-001", "title": "ยอดตามวัสดุ", "headers": ["วัสดุ", "ผลรวม", "สัดส่วน (%)"], "rows": [["ปูน", 1200, 60.0], ["เหล็ก", 800, 40.0]], "note": ""}}
        self.analysis["report"]["sections"][0]["evidence_ids"] = ["Q-001"]
        for pages in (3, 5):
            self.analysis["report"]["pages"] = pages
            self.analysis_path.write_text(json.dumps(self.analysis, ensure_ascii=False), encoding="utf-8")
            output = self.root / f"report-{pages}.pdf"
            export_report(self.database, self.analysis_path, "pdf", output)
            reader = PdfReader(output)
            self.assertEqual(len(reader.pages), pages)
            self.assertIn("1,200", reader.pages[0].extract_text() + reader.pages[1].extract_text(), "a cited result is shown as its table")


if __name__ == "__main__":
    unittest.main()
